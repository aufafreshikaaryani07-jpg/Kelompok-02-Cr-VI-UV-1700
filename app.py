"""
app.py — Media pembelajaran Spektrofotometer UV-Vis Shimadzu UV-1700 PharmaSpec
==============================================================================
Isi aplikasi:
  1. Model 3D interaktif instrumen (eksterior, cutaway, exploded view, label ID/EN,
     jalur cahaya double-beam sesuai Service Manual 206-98366A Fig. 9-1).
  2. Animasi cara kerja bertahap (sumber -> M1 -> S1 -> G -> S2 -> F -> M2 -> M3 ->
     M4/M5 -> W -> kuvet -> L -> fotodioda -> preamp -> CPU).
  3. Simulasi praktikum mengikuti IK/LAB.INST/6 (Politeknik AKA Bogor, 31 Agustus 2026).
  4. Perhitungan dengan ISIAN MANUAL: kurva kalibrasi, LOD/LOQ, residual, kadar (padat mg/kg
     atau cair), koreksi blanko, %RPD duplo, %recovery spike, diagnostik rancangan spike,
     evaluasi terhadap kriteria, dan unduhan Excel berformula. Data contoh hanya dimuat
     bila tombol "Isi contoh perhitungan" ditekan.

Menjalankan:
  pip install streamlit pandas numpy openpyxl altair
  streamlit run app.py

Model 3D memuat three.js dari CDN (cdnjs / jsDelivr), sehingga komputer perlu
terhubung ke internet saat halaman 3D dibuka.
"""
from __future__ import annotations

import io
import math
import re
from dataclasses import dataclass

import numpy as np
import pandas as pd
import streamlit as st
import streamlit.components.v1 as components

# ---------------------------------------------------------------------------
# 1. CONTOH DATA (hanya dimuat bila pengguna menekan "Isi contoh")
#    Transkripsi tampilan Sheet1 contoh Cr(VI); bukan nilai acuan.
# ---------------------------------------------------------------------------
EXAMPLE = {
    "unit": "mg/L", "mode": "abs", "basis": "padat", "zero": False,
    "std": [(0.0, 0.000), (0.1, 0.080), (0.2, 0.155), (0.3, 0.234), (0.4, 0.312), (0.5, 0.380)],
    "blank": {"use": True, "y": -0.001, "corr": False},
    "dup": [{"kode": "Sampel 1", "m": 0.4, "v": 250.0, "fp": 100.0, "y": 0.191},
            {"kode": "Sampel 2", "m": 0.4, "v": 250.0, "fp": 100.0, "y": 0.184}],
    "spk": {"kode": "Sampel 3 + spike", "m": 0.4, "v": 250.0, "fp": 100.0, "y": 0.193,
            "cstd": 1000.0, "vstd": 1.0, "base": 0.4},
    "extra": [],
}
EMPTY = {
    "unit": "mg/L", "mode": "abs", "basis": "padat", "zero": False,
    "std": [(None, None)] * 5,
    "blank": {"use": False, "y": None, "corr": False},
    "dup": [{"kode": "Duplo 1", "m": None, "v": None, "fp": 1.0, "y": None},
            {"kode": "Duplo 2", "m": None, "v": None, "fp": 1.0, "y": None}],
    "spk": {"kode": "Sampel + spike", "m": None, "v": None, "fp": 1.0, "y": None,
            "cstd": None, "vstd": None, "base": None},
    "extra": [],
}

# Kriteria eksternal yang dapat diverifikasi (dapat diedit pengguna)
DEFAULT_CRIT = pd.DataFrame({
    "Parameter": ["r kurva kalibrasi", "%RPD duplo", "%Recovery (EPA matrix spike)", "%Recovery (AOAC, tingkat ~1%)",
                  "Rasio C1/C2 (validitas spike)"],
    "Batas bawah": [0.995, 0.0, 75.0, 92.0, 0.0],
    "Batas atas": [1.0, 20.0, 125.0, 105.0, 4.0],
    "Acuan": ["Ambang linearitas umum (bukan uji linearitas)",
              "Ohio EPA Checklist 9: RPD ≤ 20% bila hasil ≥ 5×QL",
              "EPA National Functional Guidelines / Ohio EPA: 75–125%",
              "AOAC SLV (dikutip AAFCO): 92–105% pada tingkat 1%",
              "Pedoman EPA: recovery tidak dievaluasi bila kadar ≥ 4× spike"],
    "Batas lab (isi, mis. 80–120)": ["", "", "", "", ""],
})


def _ok(*v):
    return all(x is not None and not (isinstance(x, float) and math.isnan(x)) for x in v)


# ---------------------------------------------------------------------------
# 2. PERHITUNGAN (fungsi murni; tidak memakai hasil yang dilaporkan)
# ---------------------------------------------------------------------------
@dataclass
class Calib:
    b: float
    a: float
    r: float
    r2: float
    syx: float
    lod: float
    loq: float
    n: int
    zero: bool
    xmin: float
    xmax: float
    ymin: float
    ymax: float


def calibrate(pts, through_zero=False) -> Calib:
    pts = [(float(c), float(a)) for c, a in pts if _ok(c, a)]
    if len(pts) < 3:
        raise ValueError("isi minimal 3 titik standar (konsentrasi dan absorbansi).")
    x = np.array([p[0] for p in pts]); y = np.array([p[1] for p in pts])
    if np.ptp(x) == 0:
        raise ValueError("konsentrasi standar tidak boleh sama semua.")
    n = len(x)
    if through_zero:
        b = float((x * y).sum() / (x * x).sum()); a = 0.0; dof = n - 1
    else:
        b, a = (float(v) for v in np.polyfit(x, y, 1)); dof = n - 2
    if b == 0:
        raise ValueError("slope bernilai nol; periksa data standar.")
    r = float(np.corrcoef(x, y)[0, 1]) if np.ptp(y) else float("nan")
    res = y - (a + b * x)
    syx = float(math.sqrt((res ** 2).sum() / dof)) if dof > 0 else float("nan")
    return Calib(b, a, r, r * r, syx, 3.3 * syx / abs(b), 10 * syx / abs(b), n, through_zero,
                 float(x.min()), float(x.max()), float(y.min()), float(y.max()))


def residual_table(pts, c: Calib) -> pd.DataFrame:
    rows = [(float(x), float(y)) for x, y in pts if _ok(x, y)]
    d = pd.DataFrame(rows, columns=["C standar", "A terukur"])
    d["A prediksi"] = c.a + c.b * d["C standar"]
    d["Residual"] = d["A terukur"] - d["A prediksi"]
    d["C balik"] = (d["A terukur"] - c.a) / c.b
    d["Akurasi balik (%)"] = np.where(d["C standar"] != 0, d["C balik"] / d["C standar"].replace(0, np.nan) * 100, np.nan)
    return d


def kadar(c_meas, fp, v_ml, m_g, basis):
    """padat: mg/L × FP × mL / g = mg/kg ; cair: mg/L × FP = mg/L"""
    if basis == "padat":
        return c_meas * fp * v_ml / m_g
    return c_meas * fp


def rpd(x1, x2):
    m = (x1 + x2) / 2
    return abs(x1 - x2) / m * 100 if m else float("nan")


def recovery(c1, c3, c2):
    return (c3 - c1) / c2 * 100 if c2 else float("nan")


def compute(inp: dict) -> dict:
    """Menghitung seluruh hasil dari isian manual. Mengembalikan dict berisi hasil,
    langkah substitusi (untuk ditampilkan), dan peringatan."""
    out = {"warn": [], "steps": []}
    U = inp["unit"]
    K = "mg/kg" if inp["basis"] == "padat" else U
    out["K"] = K
    c = None
    if inp["mode"] == "abs":
        c = calibrate(inp["std"], inp["zero"])
        out["calib"] = c
        out["res"] = residual_table(inp["std"], c)

    def conc(y):
        return (y - c.a) / c.b if inp["mode"] == "abs" else y

    # blanko
    c_blank = 0.0
    if inp["blank"]["use"] and _ok(inp["blank"]["y"]):
        c_blank = conc(inp["blank"]["y"])
        out["c_blank"] = c_blank
        if c and c_blank > c.lod:
            out["warn"].append(f"Blanko sampel ({c_blank:.4f} {U}) di atas LOD: periksa kontaminasi pereaksi/alat.")
    corr = c_blank if inp["blank"]["corr"] else 0.0

    rows = []
    items = [("Duplo", d) for d in inp["dup"]] + [("Spike", inp["spk"])] + [("Sampel", e) for e in inp["extra"]]
    for jenis, s in items:
        need = [s.get("y"), s.get("fp")] + ([s.get("m"), s.get("v")] if inp["basis"] == "padat" else [])
        if not _ok(*need):
            continue
        cm = conc(s["y"]); cc = cm - corr
        k = kadar(cc, s["fp"], s.get("v"), s.get("m"), inp["basis"])
        rows.append({"Kode": s["kode"], "Jenis": jenis, "Hasil ukur": s["y"], "C ukur": cm, "C terkoreksi": cc,
                     "FP": s["fp"], "V (mL)": s.get("v"), "m (g)": s.get("m"), f"Kadar ({K})": k})
        if c:
            if not (c.ymin <= s["y"] <= c.ymax):
                out["warn"].append(f"{s['kode']}: absorbansi {s['y']} di luar rentang standar ({c.ymin}–{c.ymax}); encerkan atau perluas kurva.")
            if cc < c.loq:
                out["warn"].append(f"{s['kode']}: konsentrasi larutan ukur {cc:.4f} {U} di bawah LOQ ({c.loq:.4f}).")
    tab = pd.DataFrame(rows)
    out["tab"] = tab
    if tab.empty:
        return out
    kc = f"Kadar ({K})"
    dup = tab[tab["Jenis"] == "Duplo"]
    if len(dup) == 2:
        x1, x2 = dup[kc].iloc
        out.update(x1=x1, x2=x2, c1=(x1 + x2) / 2, rpd=rpd(x1, x2))
    spk = tab[tab["Jenis"] == "Spike"]
    sp = inp["spk"]
    if len(spk) and "c1" in out and _ok(sp.get("cstd"), sp.get("vstd"), sp.get("base")) and sp["base"]:
        c2 = sp["cstd"] * sp["vstd"] / sp["base"]
        c3 = spk[kc].iloc[0]
        out.update(c2=c2, c3=c3, rec=recovery(out["c1"], c3, c2))
        s = abs(out["x1"] - out["x2"]) / math.sqrt(2)
        u = s * math.sqrt(1.5)
        dg = {"Rasio C1/C2": out["c1"] / c2, "s duplo": s, "u(C3 − C1)": u, "u(%Recovery), k = 1": u / c2 * 100}
        dg["(100 − %Rec)/u"] = (100 - out["rec"]) / dg["u(%Recovery), k = 1"] if dg["u(%Recovery), k = 1"] else float("nan")
        if c:
            # kenaikan konsentrasi larutan ukur yang diharapkan dari spike
            if inp["basis"] == "padat":
                dC = c2 * sp["m"] / (sp["v"] * sp["fp"])
            else:
                dC = c2 / sp["fp"]
            dg["ΔC diharapkan di larutan ukur"] = dC
            dg["ΔA diharapkan"] = dC * c.b
            dg["ΔA teramati (A spike − rerata A duplo)"] = sp["y"] - np.mean([d["y"] for d in inp["dup"]])
            dg["ΔA teramati/diharapkan (%)"] = dg["ΔA teramati (A spike − rerata A duplo)"] / dg["ΔA diharapkan"] * 100
            if c.ymax and sp["y"] > c.ymax:
                out["warn"].append("Absorbansi spike melebihi standar tertinggi.")
        out["diag"] = dg
        if dg["Rasio C1/C2"] >= 4:
            out["warn"].append(f"Kadar asli {dg['Rasio C1/C2']:.1f}× jumlah spike (≥ 4): recovery sangat dipengaruhi presisi; "
                               "naikkan spike ke 50–150% kadar asli dengan absorbansi tetap ≤ standar tertinggi.")
        if dg["Rasio C1/C2"] < 0.2:
            out["warn"].append("Spike lebih dari 5× kadar asli: recovery tidak mewakili kondisi matriks pada kadar sebenarnya.")
    return out


def evaluate(out: dict, crit: pd.DataFrame) -> pd.DataFrame:
    vals = [out.get("calib").r if out.get("calib") else None, out.get("rpd"), out.get("rec"), out.get("rec"),
            out.get("diag", {}).get("Rasio C1/C2")]
    rows = []
    for (_, k), v in zip(crit.iterrows(), vals):
        if not _ok(v):
            se = "Data belum lengkap"
        else:
            se = "Memenuhi" if k["Batas bawah"] <= v <= k["Batas atas"] else "Tidak memenuhi"
        lab = str(k["Batas lab (isi, mis. 80–120)"] or "").strip()
        sl = "Perlu dikonfirmasi"
        mm = re.match(r"^\s*([-\d.,]+)\s*[–-]\s*([-\d.,]+)\s*$", lab)
        if mm and _ok(v):
            lo, hi = (float(g.replace(",", ".")) for g in mm.groups())
            sl = "Memenuhi" if lo <= v <= hi else "Tidak memenuhi"
        rows.append({"Parameter": k["Parameter"], "Nilai": v, "Batas acuan": f'{k["Batas bawah"]:g}–{k["Batas atas"]:g}',
                     "Acuan": k["Acuan"], "Status (acuan)": se, "Batas lab": lab or "—", "Status (lab)": sl})
    return pd.DataFrame(rows)


# ---------------------------------------------------------------------------
# 3. EXCEL berformula dari isian manual
# ---------------------------------------------------------------------------
def build_workbook(inp: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    F = "Arial"
    blue, green, blk, bold = Font(name=F, color="0000FF"), Font(name=F, color="008000"), Font(name=F), Font(name=F, bold=True)
    hdr = PatternFill("solid", fgColor="DCE6E4")
    thin = Side(style="thin", color="A6B3AF"); bd = Border(top=thin, bottom=thin, left=thin, right=thin)
    wb = Workbook(); w = wb.active; w.title = "Input"
    U = inp["unit"]; padat = inp["basis"] == "padat"; K = "mg/kg" if padat else U

    def cell(ws, r, c, v, font=blk, fmt=None, h=False):
        x = ws.cell(r, c, v if _ok(v) else None); x.font = bold if h else font; x.border = bd
        if h: x.fill = hdr; x.alignment = Alignment(wrap_text=True, horizontal="center")
        if fmt: x.number_format = fmt
        return x

    w["A1"] = "INPUT (teks biru = isian; ubah di sini, lembar Perhitungan mengikuti)"; w["A1"].font = Font(name=F, bold=True, size=12)
    meta = [("Satuan konsentrasi", U), ("Mode hasil ukur", "Absorbansi" if inp["mode"] == "abs" else "Konsentrasi"),
            ("Basis sampel", "Padat (mg/kg)" if padat else f"Cair ({U})"), ("Kurva melalui nol (1 = ya)", int(inp["zero"])),
            ("Koreksi blanko (1 = ya)", int(inp["blank"]["corr"] and inp["blank"]["use"]))]
    for i, (a, b) in enumerate(meta):
        cell(w, 3 + i, 1, a, h=True); cell(w, 3 + i, 2, b, blue)
    std = [(c, a) for c, a in inp["std"] if _ok(c, a)]
    r0 = 10
    cell(w, r0, 1, f"C standar ({U})", h=True); cell(w, r0, 2, "Absorbansi", h=True)
    for i, (c_, a_) in enumerate(std):
        cell(w, r0 + 1 + i, 1, c_, blue, "0.0000"); cell(w, r0 + 1 + i, 2, a_, blue, "0.0000")
    n = max(len(std), 1); X = f"Input!$A${r0 + 1}:$A${r0 + n}"; Y = f"Input!$B${r0 + 1}:$B${r0 + n}"
    rs = r0 + n + 3
    heads = ["Kode", "Jenis", "Bobot m (g)", "V labu (mL)", "FP", "Hasil ukur (A atau C)"]
    for j, h in enumerate(heads, 1):
        cell(w, rs, j, h, h=True)
    smp = []
    if inp["blank"]["use"] and _ok(inp["blank"]["y"]):
        smp.append(("Blanko sampel", "Blanko", None, None, None, inp["blank"]["y"]))
    for d in inp["dup"]:
        smp.append((d["kode"], "Duplo", d.get("m"), d.get("v"), d.get("fp"), d.get("y")))
    s = inp["spk"]; smp.append((s["kode"], "Spike", s.get("m"), s.get("v"), s.get("fp"), s.get("y")))
    for e in inp["extra"]:
        smp.append((e["kode"], "Sampel", e.get("m"), e.get("v"), e.get("fp"), e.get("y")))
    for i, row in enumerate(smp):
        for j, v in enumerate(row, 1):
            cell(w, rs + 1 + i, j, v, blue if j > 2 else blk, [None, None, "0.0000", "0.00", "0.###", "0.0000"][j - 1])
    rp = rs + len(smp) + 3
    spk_lab = [("C standar spike (" + U + ")", s.get("cstd")), ("V standar spike (mL)", s.get("vstd")),
               ("Bobot sampel spike (g)" if padat else "V sampel/total setelah spike (mL)", s.get("base"))]
    for i, (a, b) in enumerate(spk_lab):
        cell(w, rp + i, 1, a, h=True); cell(w, rp + i, 2, b, blue)
    for col, wd in zip("ABCDEF", [30, 16, 14, 14, 10, 20]):
        w.column_dimensions[col].width = wd

    p = wb.create_sheet("Perhitungan")
    p["A1"] = "PERHITUNGAN (rumus; hijau = tautan ke Input)"; p["A1"].font = Font(name=F, bold=True, size=12)
    absmode = inp["mode"] == "abs"
    if absmode:
        if inp["zero"]:
            f = [f"=SUMPRODUCT({X},{Y})/SUMSQ({X})", "=0", f"=CORREL({X},{Y})", "=B5^2",
                 f"=SQRT(SUMPRODUCT(({Y}-B3*{X})^2)/({n}-1))"]
        else:
            f = [f"=SLOPE({Y},{X})", f"=INTERCEPT({Y},{X})", f"=CORREL({X},{Y})", f"=RSQ({Y},{X})", f"=STEYX({Y},{X})"]
        f += ["=3.3*B7/ABS(B3)", "=10*B7/ABS(B3)"]
        labs = ["Slope b", "Intersep a", "r", "R²", "s y/x", f"LOD ({U})", f"LOQ ({U})"]
        for i, (a, fx) in enumerate(zip(labs, f)):
            cell(p, 3 + i, 1, a, h=True); cell(p, 3 + i, 2, fx, blk, "0.0000")
    t0 = 12
    for j, h in enumerate(["Kode", "Jenis", "Hasil ukur", f"C ukur ({U})", f"C terkoreksi ({U})", f"Kadar ({K})"], 1):
        cell(p, t0, j, h, h=True)
    brow = None
    for i, row in enumerate(smp):
        r, src = t0 + 1 + i, rs + 1 + i
        cell(p, r, 1, f"=Input!A{src}", green); cell(p, r, 2, f"=Input!B{src}", green); cell(p, r, 3, f"=Input!F{src}", green, "0.0000")
        cell(p, r, 4, f"=(C{r}-$B$4)/$B$3" if absmode else f"=C{r}", blk, "0.00000")
        if row[1] == "Blanko":
            brow = r; cell(p, r, 5, "—"); continue
        corr = f"-IF(Input!$B$7=1,D{brow},0)" if brow else ""
        cell(p, r, 5, f"=D{r}{corr}", blk, "0.00000")
        kf = f"=E{r}*Input!E{src}*Input!D{src}/Input!C{src}" if padat else f"=E{r}*Input!E{src}"
        cell(p, r, 6, kf, bold, "#,##0.0000")
    rowsof = lambda jen: [t0 + 1 + i for i, rr in enumerate(smp) if rr[1] == jen]
    d1, d2 = rowsof("Duplo")[:2]; sk = rowsof("Spike")[0]
    r = t0 + len(smp) + 3
    cell(p, r, 1, "%RPD = |x1 − x2| / ((x1 + x2)/2) × 100", h=True); r += 1
    cell(p, r, 1, f"C1 = rerata duplo ({K})"); cell(p, r, 2, f"=AVERAGE(F{d1},F{d2})", blk, "#,##0.0000"); rc1 = r; r += 1
    cell(p, r, 1, "%RPD"); cell(p, r, 2, f'=IFERROR(ABS(F{d1}-F{d2})/B{rc1}*100,"rerata nol")', bold, "0.00"); r += 2
    cell(p, r, 1, "%Recovery = (C3 − C1)/C2 × 100", h=True); r += 1
    cell(p, r, 1, f"C2 = C std × V std / {'m' if padat else 'V'} ({K})")
    cell(p, r, 2, f"=Input!B{rp}*Input!B{rp + 1}/Input!B{rp + 2}", blk, "#,##0.0000"); rc2 = r; r += 1
    cell(p, r, 1, f"C3 = kadar sampel + spike ({K})"); cell(p, r, 2, f"=F{sk}", blk, "#,##0.0000"); rc3 = r; r += 1
    cell(p, r, 1, "%Recovery"); cell(p, r, 2, f"=IFERROR((B{rc3}-B{rc1})/B{rc2}*100,\"\")", bold, "0.00"); r += 1
    cell(p, r, 1, "Rasio C1/C2"); cell(p, r, 2, f"=IFERROR(B{rc1}/B{rc2},\"\")", blk, "0.00")
    for col, wd in zip("ABCDEF", [38, 16, 14, 16, 18, 18]):
        p.column_dimensions[col].width = wd
    bio = io.BytesIO(); wb.save(bio)
    return bio.getvalue()


# ---------------------------------------------------------------------------
# 4. HALAMAN 3D (HTML/three.js tertanam)
# ---------------------------------------------------------------------------
HTML_3D = r"""<title>Model 3D UV-1700 PharmaSpec</title>
<meta name="claude:viewport-fit-sides" content="cover">
<link rel="preconnect" href="https://fonts.googleapis.com">
<link rel="stylesheet" href="https://fonts.googleapis.com/css2?family=IBM+Plex+Sans:wght@400;500;600&family=IBM+Plex+Mono:wght@400;500&family=Fraunces:opsz,wght@9..144,560&display=swap">
<style>
/* Layout: instrument viewport on top (bench lamp lighting), working panel with tabs below; two columns from 1100px */
:root{
  --bg:#eef0ed; --panel:#ffffff; --ink:#1d2421; --muted:#5d6762; --line:#d3d9d5;
  --casing:#d8c48c; --accent:#0f6e74; --accent-ink:#ffffff; --warn:#9a5b00; --ok:#1f7a3a; --bad:#a3292b;
  --uv:#8b5cf6; --vis:#e0a21b; --sam:#13a89e; --ref:#ea6a2e; --stage:#121615; --stage-ink:#e9eee9;
  --f-display:"Fraunces",Georgia,serif; --f-body:"IBM Plex Sans",system-ui,sans-serif; --f-mono:"IBM Plex Mono",ui-monospace,monospace;
}
@media (prefers-color-scheme:dark){:root:not([data-theme="light"]){
  --bg:#121514; --panel:#1b201e; --ink:#e6ebe8; --muted:#9aa49f; --line:#2f3733; --accent:#4cc2c4; --accent-ink:#06201f;
  --warn:#e6a743; --ok:#5fc27d; --bad:#ef7a7c; color-scheme:dark}}
:root[data-theme="dark"]{--bg:#121514; --panel:#1b201e; --ink:#e6ebe8; --muted:#9aa49f; --line:#2f3733; --accent:#4cc2c4; --accent-ink:#06201f;
  --warn:#e6a743; --ok:#5fc27d; --bad:#ef7a7c; color-scheme:dark}
html,body{background:var(--bg);color:var(--ink)}
body{font-family:var(--f-body);font-size:15px;line-height:1.55;padding-left:max(16px,env(safe-area-inset-left,0px));padding-right:max(16px,env(safe-area-inset-right,0px));padding-block:14px 40px}
*{box-sizing:border-box}
h1,h2,h3{text-wrap:balance;margin:0}
h1{font-family:var(--f-display);font-weight:560;font-size:clamp(1.35rem,3.6vw,1.9rem);letter-spacing:-.01em}
h2{font-size:1.05rem;font-weight:600;margin-bottom:.4rem}
h3{font-size:.95rem;font-weight:600}
p{margin:.35rem 0;max-width:70ch}
.mono{font-family:var(--f-mono);font-variant-numeric:tabular-nums}
.eyebrow{font-family:var(--f-mono);font-size:.72rem;letter-spacing:.08em;text-transform:uppercase;color:var(--muted)}
header{display:flex;flex-wrap:wrap;align-items:baseline;justify-content:space-between;gap:6px 16px;margin-bottom:10px}
.wrap{display:grid;gap:14px;max-width:1500px;margin:0 auto}
@media(min-width:1100px){.wrap{grid-template-columns:minmax(0,1.35fr) minmax(0,1fr);align-items:start}.side{position:sticky;top:12px;max-height:calc(100vh - 24px);overflow:auto}}
.stagebox{background:var(--stage);border-radius:10px;overflow:hidden;position:relative;min-width:0}
#stage{position:relative;height:min(62vh,620px);min-height:340px;touch-action:none}
#stage canvas{display:block;width:100%;height:100%}
#labels{position:absolute;inset:0;pointer-events:none;overflow:hidden}
#labels svg{position:absolute;inset:0;width:100%;height:100%}
.lbl{position:absolute;pointer-events:auto;cursor:pointer;background:rgba(18,22,21,.86);color:var(--stage-ink);border:1px solid rgba(255,255,255,.18);border-radius:4px;padding:2px 6px;font-size:11px;line-height:1.25;white-space:nowrap;transform:translate(-50%,-50%)}
.lbl b{font-weight:600;display:block}.lbl i{font-style:normal;opacity:.72;display:block;font-size:10px}
.lbl .c{font-family:var(--f-mono);color:#9fe3e0;margin-right:4px}
.lbl.sel{border-color:#9fe3e0;background:#0d3c3f}
.toolbar{display:flex;flex-wrap:wrap;gap:6px;padding:10px;background:#1a1f1d;border-top:1px solid #2b3330}
.tb{font:500 12px var(--f-body);color:var(--stage-ink);background:#262d2a;border:1px solid #39433f;border-radius:6px;padding:6px 10px;cursor:pointer}
.tb[aria-pressed="true"]{background:#0f6e74;border-color:#3fb6b8}
.tb:focus-visible,button:focus-visible,select:focus-visible,input:focus-visible{outline:2px solid #4cc2c4;outline-offset:2px}
.ctl{display:flex;flex-wrap:wrap;gap:10px 18px;align-items:center;padding:10px;background:#1a1f1d;color:var(--stage-ink);font-size:12px;border-top:1px solid #2b3330}
.ctl label{display:flex;align-items:center;gap:8px}
.ctl input[type=range]{width:min(260px,52vw);accent-color:#4cc2c4}
.ctl .v{font-family:var(--f-mono);font-variant-numeric:tabular-nums;min-width:5.5ch;color:#fff}
.legend{display:flex;flex-wrap:wrap;gap:8px 14px;padding:8px 10px;background:#141917;color:#c9d2cd;font-size:11px}
.legend span{display:inline-flex;align-items:center;gap:6px}.legend i{width:18px;height:4px;border-radius:2px;display:inline-block}
#caption{position:absolute;left:10px;right:10px;bottom:10px;background:rgba(12,16,15,.9);color:var(--stage-ink);border:1px solid #33413c;border-radius:8px;padding:8px 10px;font-size:12.5px;display:none}
#caption b{color:#9fe3e0}
#hint{position:absolute;top:8px;left:10px;color:#a8b3ae;font-size:11px;pointer-events:none}
.panel{background:var(--panel);border:1px solid var(--line);border-radius:10px;padding:14px;min-width:0}
#info .code{font-family:var(--f-mono);color:var(--accent);font-size:.8rem}
.pill{display:inline-block;font:500 11px var(--f-mono);padding:1px 7px;border-radius:999px;border:1px solid currentColor;margin-right:6px}
.pill.v{color:var(--ok)}.pill.k{color:var(--warn)}.pill.x{color:var(--bad)}
.tabs{display:flex;gap:2px;flex-wrap:wrap;border-bottom:1px solid var(--line);margin:-4px -4px 12px}
.tab{font:500 13px var(--f-body);background:none;border:0;color:var(--muted);padding:8px 10px;cursor:pointer;border-bottom:2px solid transparent}
.tab[aria-selected="true"]{color:var(--ink);border-bottom-color:var(--accent)}
.tblwrap{overflow-x:auto;margin:6px 0}
table{border-collapse:collapse;width:100%;font-size:13px}
th,td{text-align:left;padding:6px 8px;border-bottom:1px solid var(--line);vertical-align:top}
th{font-weight:600;font-size:12px;color:var(--muted);white-space:nowrap}
td.n,th.n{text-align:right;font-family:var(--f-mono);font-variant-numeric:tabular-nums;white-space:nowrap}
tbody tr.clk{cursor:pointer}tbody tr.clk:hover{background:color-mix(in srgb,var(--accent) 8%,transparent)}
.flow{display:flex;flex-wrap:wrap;gap:6px;align-items:center;margin:8px 0}
.flow button{font:500 12px var(--f-body);background:var(--bg);color:var(--ink);border:1px solid var(--line);border-radius:6px;padding:5px 8px;cursor:pointer}
.flow button.on{background:var(--accent);color:var(--accent-ink);border-color:var(--accent)}
.flow .ar{color:var(--muted);font-size:12px}
.eq{font-family:var(--f-mono);background:var(--bg);border:1px solid var(--line);border-radius:6px;padding:8px 10px;margin:6px 0;overflow-x:auto;white-space:nowrap}
.note{border-left:3px solid var(--warn);padding:4px 10px;margin:8px 0;color:var(--ink);background:color-mix(in srgb,var(--warn) 8%,transparent)}
.safe{border-left:3px solid var(--bad);padding:4px 10px;margin:8px 0;background:color-mix(in srgb,var(--bad) 7%,transparent)}
ul,ol{padding-left:1.2rem;margin:.3rem 0}li{margin:.15rem 0}
.btn{font:600 13px var(--f-body);background:var(--accent);color:var(--accent-ink);border:0;border-radius:6px;padding:8px 12px;cursor:pointer}
.btn.sec{background:transparent;color:var(--accent);border:1px solid var(--accent)}
.btn:disabled{opacity:.45;cursor:not-allowed}
.row{display:flex;flex-wrap:wrap;gap:8px;align-items:center}
select,input[type=number],input[type=text]{font:13px var(--f-body);background:var(--bg);color:var(--ink);border:1px solid var(--line);border-radius:6px;padding:6px 8px}
.steps{list-style:none;padding:0;margin:8px 0;max-height:260px;overflow:auto;border:1px solid var(--line);border-radius:8px}
.steps li{padding:6px 10px;border-bottom:1px solid var(--line);font-size:13px;display:flex;gap:8px}
.steps li:last-child{border-bottom:0}
.steps li .s{font-family:var(--f-mono);color:var(--muted);min-width:3.2ch}
.steps li.done{color:var(--muted)}.steps li.done .s{color:var(--ok)}
.steps li.cur{background:color-mix(in srgb,var(--accent) 10%,transparent);font-weight:500}
.simbar{display:flex;flex-wrap:wrap;gap:6px;align-items:center;margin:6px 0}.seg{display:inline-flex;border:1px solid var(--line);border-radius:6px;overflow:hidden}.spd{font:600 12px var(--f-mono);background:var(--bg);color:var(--ink);border:0;padding:7px 10px;cursor:pointer}.spd[aria-pressed="true"]{background:var(--accent);color:var(--accent-ink)}.prog{height:4px;background:var(--line);border-radius:2px;overflow:hidden;margin-bottom:8px}#simbar{height:100%;width:0;background:var(--accent)}#simcap{position:absolute;left:10px;right:10px;top:8px;background:rgba(12,16,15,.9);color:var(--stage-ink);border:1px solid #33413c;border-radius:8px;padding:7px 10px;font-size:12.5px;display:none;z-index:2}#simcap i{opacity:.75}#simcap .menu{background:#26302c;border-color:#3a4842;color:#dfe9e4}
.cur-step{border:1px solid var(--accent);border-radius:8px;padding:10px 12px;margin:8px 0}
.menu{font-family:var(--f-mono);font-size:12px;background:var(--bg);border:1px solid var(--line);border-radius:4px;padding:1px 5px}
.log{font-family:var(--f-mono);font-size:12px;background:var(--bg);border:1px solid var(--line);border-radius:6px;padding:8px;max-height:150px;overflow:auto;white-space:pre-wrap}
.grid2{display:grid;gap:12px}@media(min-width:640px){.grid2{grid-template-columns:1fr 1fr}}
.kv{display:grid;grid-template-columns:auto 1fr;gap:2px 12px;font-size:13px}.kv dt{color:var(--muted)}.kv dd{margin:0;font-family:var(--f-mono);font-variant-numeric:tabular-nums}
svg text{font-family:var(--f-mono)}
@media (prefers-reduced-motion:reduce){*{transition:none!important}}
</style>

<header>
  <div>
    <h1>Spektrofotometer UV-Vis UV-1700 PharmaSpec — sistem optik double-beam</h1>
  </div>
  <div class="eyebrow">Susunan optik: Service Manual 206-98366A, Fig. 9-1 · Prosedur: IK/LAB.INST/6</div>
</header>

<div class="wrap">
  <div class="stagebox">
    <div id="stage"><div id="labels"><svg id="leaders"></svg></div><div id="hint">Seret: putar 360° · cubit/gulir: zoom · ketuk komponen untuk detail</div><div id="caption"></div><div id="simcap" aria-live="polite"></div></div>
    <div class="toolbar" role="toolbar" aria-label="Kontrol tampilan">
      <button class="tb" id="bCase" aria-pressed="false">Casing tampak</button>
      <button class="tb" id="bCut" aria-pressed="true">Cutaway</button>
      <button class="tb" id="bExp" aria-pressed="false">Exploded view</button>
      <button class="tb" id="bLbl">Label: lengkap</button>
      <button class="tb" id="bBeam" aria-pressed="true">Jalur cahaya</button>
      <button class="tb" id="bStep" aria-pressed="false">Mode bertahap</button>
      <button class="tb" id="bPrev" hidden>◀</button><button class="tb" id="bNext" hidden>▶</button>
      <button class="tb" id="bReset">Reset kamera</button>
    </div>
    <div class="ctl">
      <label for="wl">λ terpasang <input type="range" id="wl" min="190" max="1100" step="0.5" value="540"><span class="v" id="wlv">540.0 nm</span></label>
      <label for="sw">λ pergantian lampu <input type="number" id="sw" min="295" max="364" step="0.1" value="340.8" style="width:6.5em"> nm</label>
      <span id="stat" class="mono"></span>
    </div>
    <div class="legend">
      <span><i style="background:var(--uv)"></i>UV polikromatik (D₂)</span>
      <span><i style="background:var(--vis)"></i>Tampak/NIR polikromatik (WI)</span>
      <span><i style="background:linear-gradient(90deg,#7b3fe4,#2f8cff,#2fd06b,#ffd32f,#ff4b2f)"></i>Monokromatik (warna = λ)</span>
      <span><i style="background:var(--sam)"></i>Berkas sampel (S)</span>
      <span><i style="background:var(--ref)"></i>Berkas referensi (R)</span>
    </div>
  </div>

  <div class="side" style="display:grid;gap:14px;min-width:0">
    <section class="panel" id="info" aria-live="polite"></section>
    <section class="panel">
      <div class="tabs" role="tablist">
        <button class="tab" role="tab" aria-selected="true" data-t="tKomp">Komponen</button>
        <button class="tab" role="tab" aria-selected="false" data-t="tPrin">Prinsip &amp; alur</button>
        <button class="tab" role="tab" aria-selected="false" data-t="tPrak">Simulasi praktikum</button>
        <button class="tab" role="tab" aria-selected="false" data-t="tData">Data Cr(VI)</button>
        <button class="tab" role="tab" aria-selected="false" data-t="tSrc">Sumber</button>
      </div>

      <div id="tKomp" role="tabpanel">
        <p>Status <span class="pill v">Terverifikasi</span> berarti komponen dan urutannya tercantum dalam Service Manual. <span class="pill k">Inferensi/konseptual</span> berarti bentuk, posisi, atau gerak digambarkan untuk keperluan didaktik dan tidak dinyatakan secara eksplisit dalam dokumen. Jarak dan sudut seluruh komponen bersifat skematik; yang dijaga adalah topologi jalur optik.</p>
        <div class="tblwrap"><table id="ktab"><thead><tr><th>Kode</th><th>Komponen (ID / EN)</th><th>Fungsi</th><th>Status</th><th>Sumber</th></tr></thead><tbody></tbody></table></div>
      </div>

      <div id="tPrin" role="tabpanel" hidden>
        <h2>Prinsip</h2>
        <p>Spektrofotometri UV-Vis mengukur fraksi radiasi monokromatik 190–1100 nm yang diserap analit akibat transisi elektronik, lalu mengonversinya menjadi absorbansi yang sebanding dengan konsentrasi.</p>
        <div class="eq">A = log₁₀(I₀/I) = ε·b·c</div>
        <p class="mono" style="font-size:12.5px">A absorbansi (tanpa satuan) · I₀ intensitas acuan · I intensitas yang diteruskan · ε absorptivitas molar (L·mol⁻¹·cm⁻¹) · b panjang lintasan (cm; kuvet standar 1 cm) · c konsentrasi (mol·L⁻¹)</p>
        <p>Linearitas mensyaratkan radiasi monokromatik, larutan encer tanpa interaksi antarmolekul, serta tidak adanya hamburan, fluoresensi, atau cahaya liar. Cahaya liar (spesifikasi ≤0,04% pada 220 dan 340 nm) menjadi pembatas utama pada absorbansi tinggi karena menambah I semu sehingga kurva kalibrasi melengkung ke bawah.</p>

        <h2 style="margin-top:12px">Alur kerja (ketuk untuk menyorot bagian 3D)</h2>
        <div class="flow" id="flow"></div>
        <div class="note" id="flowtxt">Pilih satu tahap.</div>

        <h3 style="margin-top:10px">Pemilihan panjang gelombang</h3>
        <p>Cahaya dari lampu aktif dipantulkan M1 ke celah masuk S1. Kisi difraksi G mendispersi cahaya menurut persamaan kisi <span class="mono">mλ = d(sin α + sin β)</span>; rotasi G (penggerak Harmonic Drive RH-8-SP) menentukan λ yang jatuh tepat pada celah keluar S2. Lebar celah menetapkan lebar pita spektral 1 nm. Kisi juga meloloskan orde dua dan tiga (λ/2, λ/3) ke S2, sehingga filter F berpindah otomatis ke salah satu dari enam posisi untuk menahan orde lebih tinggi tersebut.</p>
        <div class="tblwrap"><table><thead><tr><th>Filter</th><th class="n">Rentang λ (nm)</th></tr></thead><tbody id="ftab"></tbody></table></div>
        <h3 style="margin-top:10px">D₂ dan halogen (WI)</h3>
        <p>Lampu deuterium menghasilkan kontinum UV melalui lucutan busur dalam gas D₂; lampu halogen tungsten-iodin menghasilkan spektrum benda pijar yang dominan di daerah tampak dan NIR. Kedua spektrum berpotongan di daerah UV dekat; titik pergantian dapat diatur 295,0–364,0 nm (bawaan 340,8 nm), dan M1 berputar otomatis memilih lampu sesuai λ terpasang. Umur rata-rata: WI 2000 jam, D₂ 500 jam.</p>
        <h3 style="margin-top:10px">Pembentukan berkas S dan R serta hasil fotometrik</h3>
        <p>Setelah filter, cermin toroidal M2 mengarahkan berkas ke pembagi berkas M3 (half mirror) yang membaginya menjadi berkas sampel (via M4) dan referensi (via M5). Kedua berkas melewati jendela kuarsa W, terfokus di pusat pemegang kuvet, dikumpulkan lensa L, dan jatuh pada dua fotodioda Si. Preamplifier mengubah arus foto menjadi tegangan; CPU melakukan konversi A/D dan pengolahan sinyal.</p>
        <p>Secara prinsip, absorbansi diperoleh dari rasio intensitas kedua kanal: <span class="mono">A ≈ log₁₀(I_R/I_S) − A_baseline(λ)</span>. Koreksi baseline atau Autozero dengan blanko di kedua posisi menghapus ketidaksamaan respons kedua kanal. Pengukuran simultan S dan R mengompensasi fluktuasi intensitas lampu, tetapi tidak mengompensasi perbedaan kuvet, endapan pada kuvet, atau drift termal detektor; karena itu SOP mensyaratkan pemanasan 15 menit. Algoritme internal tidak dirinci dalam dokumen yang tersedia; persamaan di atas adalah model konseptual.</p>
      </div>

      <div id="tPrak" role="tabpanel" hidden>
        <p>Urutan dan nama menu mengikuti <b>IK/LAB.INST/6</b> (Politeknik AKA Bogor, berlaku 31 Agustus 2026). Nilai absorbansi yang muncul saat pembacaan diambil dari data praktikum Anda; tidak ada angka yang dibuat. Langkah berjalan otomatis; waktu pemanasan dipercepat.</p>
        <div class="simbar" role="toolbar" aria-label="Kontrol simulasi">
          <button class="btn sec" id="simrst" title="Ulang dari awal">⟲ Ulang</button>
          <button class="btn sec" id="simprev">◀ Sebelumnya</button>
          <button class="btn" id="simplay" aria-pressed="false">▶ Putar</button>
          <button class="btn sec" id="simnext">Berikutnya ▶</button>
          <span class="seg" role="group" aria-label="Kecepatan"><button class="spd" data-speed="0.5" aria-pressed="false">0,5×</button><button class="spd" data-speed="1" aria-pressed="true">1×</button><button class="spd" data-speed="2" aria-pressed="false">2×</button></span>
        </div>
        <div class="row" style="margin:6px 0;font-size:13px">
          <label for="simlang">Bahasa prosedur <select id="simlang"><option value="both" selected>Indonesia + English</option><option value="id">Indonesia</option><option value="en">English</option></select></label>
          <label><input type="checkbox" id="simzoom" checked> Zoom otomatis ke bagian yang bekerja</label>
          <label for="simwl">λ maks (nm) <input type="number" id="simwl" value="540" step="0.1" style="width:6em"></label>
        </div>
        <div class="cur-step" id="simcur"><div class="eyebrow">Simulasi siap</div><p style="margin:4px 0">Tekan Putar. Langkah berjalan otomatis mengikuti IK; kamera mendekat ke bagian alat yang sedang bekerja.</p></div>
        <div class="prog"><div id="simbar"></div></div>
        <div class="grid2">
          <div><h3>Kompartemen sampel</h3><p class="mono" id="slots" style="font-size:12px"></p>
            <svg id="cal" viewBox="0 0 360 170" style="width:100%;height:auto" role="img" aria-label="Kurva kalibrasi yang terbentuk dari pembacaan standar"></svg></div>
          <div><h3>Catatan pembacaan</h3><div class="log" id="log"></div></div>
        </div>
        <h3 style="margin-top:10px">Seluruh langkah IK (ketuk untuk melompat)</h3><ol class="steps" id="steplist"></ol>
        <h3 style="margin-top:12px">Jenis larutan</h3>
        <ul>
          <li><b>Blanko</b>: matriks pelarut/pereaksi tanpa analit; menetapkan I₀ (baseline, Autozero). Blanko sampel (blanko prosedur) melalui seluruh preparasi untuk mendeteksi kontaminasi.</li>
          <li><b>Larutan standar</b>: konsentrasi analit diketahui dari bahan acuan; membentuk kurva kalibrasi A = a + b·C.</li>
          <li><b>Larutan sampel</b>: hasil preparasi sampel; konsentrasinya diinterpolasi dari kurva lalu dikoreksi faktor pengenceran, volume, dan bobot.</li>
          <li><b>Larutan spike</b>: porsi sampel yang ditambah standar dengan jumlah diketahui sebelum preparasi; selisih terhadap sampel asli dibandingkan dengan jumlah yang ditambahkan (%recovery) untuk menilai kehilangan atau interferensi matriks.</li>
        </ul>
        <h3 style="margin-top:12px">Keselamatan dan penanganan</h3>
        <div class="safe"><ul>
          <li>Jangan membuka casing atau menyentuh bagian internal saat alat menyala; perbaikan internal hanya oleh teknisi terlatih (prasyarat Service Manual).</li>
          <li>Rumah lampu panas (label HOT SURFACE): biarkan lampu dingin hingga suhu ruang sebelum penggantian; jangan menyentuh kaca lampu dengan tangan telanjang.</li>
          <li>Jangan menghalangi kipas pendingin di sisi kiri; kipas mati memicu sakelar suhu dan lampu WI tidak menyala.</li>
          <li>Bersihkan tumpahan di lantai kompartemen sebelum dan sesudah analisis untuk mencegah korosi oleh uap larutan.</li>
          <li>Jauhkan telepon seluler dari sisi kiri alat (dekat preamplifier) selama pengukuran; dapat menimbulkan derau.</li>
          <li>Cr(VI) bersifat karsinogenik: gunakan sarung tangan, jas lab, dan tampung limbah sesuai prosedur laboratorium.</li>
        </ul></div>
        <div class="note"><ul>
          <li>Pegang kuvet pada sisi buram; seka sisi bening dengan tisu lensa; isi ±¾ volume tanpa gelembung.</li>
          <li>Sisi bening menghadap berkas (berkas berjalan kanan → kiri melewati pemegang; perhatikan tanda panah pada pemegang).</li>
          <li>Gunakan kuvet kuarsa untuk λ &lt; ±320 nm; kaca menyerap UV.</li>
          <li>Pastikan label S/R pada pemegang di kompartemen Anda; model ini menempatkan S di depan dan R di belakang (perlu verifikasi pada alat).</li>
        </ul></div>
      </div>

      <div id="tData" role="tabpanel" hidden>
        <p>Rekapitulasi dihitung ulang dari nilai yang terlihat pada tangkapan layar Sheet1. Buku kerja Excel terpisah memuat rumus lengkap dan dapat diperbarui begitu berkas aslinya diunggah.</p>
        <div id="dataBox"></div>
      </div>

      <div id="tSrc" role="tabpanel" hidden>
        <ol id="srcs" style="font-size:13px"></ol>
      </div>
    </section>
  </div>
</div>

<script src="https://cdnjs.cloudflare.com/ajax/libs/three.js/r128/three.min.js"></script>
<script src="https://cdn.jsdelivr.net/npm/three@0.128.0/examples/js/controls/OrbitControls.js"></script>
<script>
(()=>{
const SRC={
 SM:'Service Manual UV-1700 206-98366A',
};
/* ---------- component knowledge base ---------- */
const INFO={
 d2:{code:'D₂',id:'Lampu deuterium',en:'Deuterium lamp',st:'v',src:'SM §9.1; §10.6 (D2 Lamp L6380); §5.2.2',
   fn:'Sumber radiasi kontinu daerah UV.',pr:'Lucutan busur dalam gas deuterium bertekanan rendah mengeksitasi molekul D₂; peluruhannya memancarkan kontinum UV. Digunakan di bawah λ pergantian lampu. Umur rata-rata 500 jam; garis emisi 656,1 dan 486,0 nm dipakai untuk uji akurasi panjang gelombang.'},
 wi:{code:'WI',id:'Lampu halogen (tungsten-iodin)',en:'Halogen lamp',st:'v',src:'SM §9.1; §10.6 (Halogen lamp NA55917); §5.2.2',
   fn:'Sumber radiasi kontinu daerah tampak dan NIR.',pr:'Filamen tungsten pijar memancarkan spektrum benda hitam; siklus halogen meredeposisi tungsten sehingga intensitas stabil. Digunakan di atas λ pergantian lampu. Umur rata-rata 2000 jam.'},
 m1:{code:'M1',id:'Cermin pemilih sumber cahaya',en:'Light source switch mirror',st:'v',src:'SM §9.1; §8.2.2 (Light source mirror PM); §10.6',
   fn:'Memilih lampu aktif dan memantulkan cahayanya ke celah masuk.',pr:'Berputar otomatis sesuai λ terpasang (D₂ untuk UV, WI untuk tampak/UV dekat). Titik pergantian 295,0–364,0 nm, bawaan 340,8 nm. Rotasi pada model menunjukkan pergantian arah secara konseptual; sudut sebenarnya tidak didokumentasikan.'},
 s1:{code:'S1',id:'Celah masuk',en:'Entrance slit',st:'v',src:'SM §9.1; §10.9 (Slit & Filter Assy)',
   fn:'Membatasi berkas menjadi sumber garis sempit untuk monokromator.',pr:'Lebar celah, bersama dispersi kisi, menentukan lebar pita spektral (1 nm, 190–900 nm). Celah lebih sempit meningkatkan resolusi tetapi menurunkan energi.'},
 g:{code:'G',id:'Kisi difraksi',en:'Diffraction grating',st:'v',src:'SM §9.1; §8.3.3; §10.7 (Grating Assy, Harmonic Drive RH-8-SP)',
   fn:'Mendispersi cahaya polikromatik dan memilih λ yang menuju celah keluar.',pr:'Mengikuti mλ = d(sin α + sin β). Rotasi kisi oleh motor penggerak mengubah λ yang terfokus di S2. Fig. 9-1 tidak memuat cermin kolimator di antara S1–G–S2; dari sini disimpulkan kisi bertipe cekung yang sekaligus memfokuskan (inferensi). Sudut rotasi pada model bersifat konseptual.'},
 s2:{code:'S2',id:'Celah keluar',en:'Exit slit',st:'v',src:'SM §9.1; §8.3.6',
   fn:'Meloloskan pita sempit di sekitar λ terpasang.',pr:'Hanya cahaya yang difokuskan kisi tepat pada bukaan yang keluar; panjang gelombang lain terhalang. Fokus celah keluar disetel dengan garis Hg 546,1 nm.'},
 f:{code:'F',id:'Filter pemotong cahaya liar',en:'Stray light cut filter',st:'v',src:'SM §9.1, Fig. 9-2; §10.9 (motor PFC25-24-1)',
   fn:'Menahan cahaya orde dua dan tiga (λ/2, λ/3) yang ikut lolos dari kisi.',pr:'Roda filter enam posisi berpindah otomatis menurut λ: tanpa filter (190–305), U-340 (306–364), B-390 (365–416), B-460 (417–536), Y-52 (537–760), R-72 (761–1100 nm). Filter yang rusak menyebabkan baseline tidak rata dan kegagalan uji λ.'},
 m2:{code:'M2',id:'Cermin toroidal',en:'Toroidal mirror',st:'v',src:'SM §9.1; §8.3.4; §10.7 (Toroidal Mirror Assy)',
   fn:'Mengarahkan dan membentuk berkas monokromatik menuju pembagi berkas.',pr:'Permukaan toroidal memiliki jari-jari kelengkungan berbeda pada dua sumbu sehingga dapat mengoreksi astigmatisme pada pantulan miring (penjelasan optik umum; inferensi).'},
 m3:{code:'M3',id:'Pembagi berkas (half mirror)',en:'Beam splitter',st:'v',src:'SM §9.1; §8.3.4',
   fn:'Membagi berkas menjadi berkas sampel dan berkas referensi.',pr:'Cermin setengah memantulkan sebagian dan meneruskan sebagian intensitas, sehingga kedua kanal diukur serentak (double-beam tanpa chopper berputar).'},
 m4:{code:'M4',id:'Cermin jalur sampel',en:'Sample side mirror',st:'v',src:'SM §9.1; §8.3.4 (S side mirror)',
   fn:'Membelokkan berkas sampel ke pusat pemegang kuvet sampel.',pr:'Posisi dan kemiringannya disetel agar berkas melewati pusat kuvet pada ketinggian yang sama dengan berkas R.'},
 m5:{code:'M5',id:'Cermin jalur referensi',en:'Reference side mirror',st:'v',src:'SM §9.1; §8.3.4 (R side mirror)',
   fn:'Membelokkan berkas referensi ke pusat pemegang kuvet referensi.',pr:'Disetel bersama M2 dan M4 dalam penyetelan optik double-beam.'},
 w:{code:'W',id:'Jendela kuarsa',en:'Quartz windowpane',st:'v',src:'SM §9.1; §6.2',
   fn:'Memisahkan kompartemen sampel dari ruang optik sambil meneruskan UV.',pr:'Kuarsa transparan hingga UV jauh. Jendela kotor menyebabkan baseline tidak rata.'},
 cellS:{code:'Sam.',id:'Pemegang kuvet sampel + kuvet',en:'Sample side cell holder',st:'v',src:'SM §9.1–9.2; §10.8; foto kompartemen',
   fn:'Menempatkan kuvet sampel pada lintasan berkas S.',pr:'Berkas berpusat di tengah pemegang; posisi berkas sama untuk sisi S dan R. Penempatan S di depan pada model perlu dicocokkan dengan label pada alat.'},
 cellR:{code:'Ref.',id:'Pemegang kuvet referensi + kuvet',en:'Reference side cell holder',st:'v',src:'SM §9.1–9.2; §10.8; foto kompartemen',
   fn:'Menempatkan kuvet blanko/referensi pada lintasan berkas R.',pr:'Dalam koreksi baseline kedua posisi berisi blanko; pada pengukuran rutin posisi R tetap berisi blanko.'},
 l:{code:'L',id:'Lensa',en:'Lens',st:'v',src:'SM §9.1',
   fn:'Mengumpulkan berkas yang keluar dari kuvet ke permukaan fotodioda.',pr:'Konvergensi berkas mengurangi kepekaan terhadap pergeseran berkas akibat kuvet atau pembiasan larutan.'},
 pd:{code:'P.D.',id:'Detektor fotodioda silikon (×2)',en:'Si photodiode',st:'v',src:'SM §9.1; §9.4; §10.12 (Photodiode S1336-8BQ, D1–D2)',
   fn:'Mengubah intensitas cahaya kanal S dan R menjadi arus listrik.',pr:'Foton dengan energi di atas celah pita silikon membangkitkan pasangan elektron-hole; arus foto sebanding dengan intensitas dalam rentang linear. Respons spektral Si mencakup sekitar 190–1100 nm.'},
 pre:{code:'PCB',id:'Preamplifier',en:'Preamp PCB',st:'v',src:'SM §9.4; §10.12 (op-amp AD645KN, resistor 200 MΩ); §6.3',
   fn:'Mengubah arus foto menjadi tegangan sinyal.',pr:'Penguat transimpedans masukan-FET dengan resistor umpan balik bernilai sangat tinggi (inferensi dari daftar komponen). Terletak di sisi kiri; peka terhadap gangguan elektromagnetik. Posisi pada model skematik.'},
 cpu:{code:'PCB',id:'PCB CPU',en:'CPU PCB',st:'v',src:'SM §9.4; §7.5',
   fn:'Mengendalikan sistem dan mengolah sinyal fotometrik.',pr:'Melakukan konversi A/D, perhitungan fotometrik, kendali lampu, pergantian sumber, penggerak λ, filter, papan tombol, LCD, printer, dan komunikasi RS-232C. Dicabut dari sisi kanan unit. Posisi pada model skematik.'},
 pwr:{code:'PCB',id:'PCB catu daya + transformator',en:'Power PCB, transformer',st:'k',src:'SM §9.4; §10.5; §10.11',
   fn:'Menyuplai daya sistem dan lampu halogen serta D₂.',pr:'Kebutuhan daya 130 VA; sekering 2,0 A untuk sistem 200 V. Bentuk dan posisi pada model konseptual.'},
 fan:{code:'FAN',id:'Kipas pendingin',en:'Cooling fan',st:'v',src:'SM §2.3.3; §5.1.1; §10.5; foto sisi kiri',
   fn:'Mendinginkan kompartemen sumber cahaya dan elektronik.',pr:'Terletak di sisi kiri bersama sakelar daya. Jika kipas mati, sakelar suhu aktif dan lampu WI tidak menyala.'},
 con:{code:'LCD',id:'Layar LCD dan papan tombol',en:'Display & keypad (console)',st:'v',src:'SM §10.4; foto alat',
   fn:'Antarmuka operasi mandiri dan pemilihan mode PC Control.',pr:'Tombol pada foto: F1–F4, RETURN, MODE, PRINT, angka, GOTO WL, AUTO ZERO, ENTER, START/STOP. Pada IK, F4 dipakai untuk PC Control.'},
 lid:{code:'',id:'Tutup kompartemen sampel',en:'Sample compartment cover',st:'v',src:'SM §10.4; foto alat',
   fn:'Mencegah cahaya luar masuk ke detektor selama pengukuran.',pr:'Selalu ditutup saat baseline, autozero, dan pembacaan.'},
 io:{code:'I/O',id:'Port antarmuka dan slot IC card',en:'I/O ports, IC card slot',st:'v',src:'Foto sisi kanan; SM §9.4',
   fn:'Koneksi ke PC (RS-232C), printer (Centronics), aksesori (I/O-2), dan penyimpanan IC card.',pr:'Kabel data menghubungkan alat dengan PC yang menjalankan UVProbe.'},
 psw:{code:'',id:'Sakelar daya',en:'Power switch',st:'v',src:'SM §2.3.3; IK A.3; foto sisi kiri',
   fn:'Menyalakan dan mematikan alat.',pr:'Terletak di sisi kiri. Setelah dinyalakan, alat menjalankan inisialisasi otomatis.'}
};
const FILT=[['Tanpa filter',190,305,'#3a3f45'],['U-340',306,364,'#4b2a7a'],['B-390',365,416,'#2d3f9e'],['B-460',417,536,'#1f6fb0'],['Y-52',537,760,'#d6a51c'],['R-72',761,1100,'#9c2121']];

/* ---------- three setup ---------- */
const stage=document.getElementById('stage');
const renderer=new THREE.WebGLRenderer({antialias:true});
renderer.setPixelRatio(Math.min(devicePixelRatio,2));
renderer.outputEncoding=THREE.sRGBEncoding;
stage.prepend(renderer.domElement);
const scene=new THREE.Scene();scene.background=new THREE.Color(0x121615);
const camera=new THREE.PerspectiveCamera(38,1,0.5,600);
const controls=new THREE.OrbitControls(camera,renderer.domElement);
controls.enableDamping=true;controls.minDistance=18;controls.maxDistance=170;
const HOME={p:new THREE.Vector3(48,52,66),t:new THREE.Vector3(-1,7,-2)};
function resetCam(){camera.position.copy(HOME.p);controls.target.copy(HOME.t);controls.update();}
resetCam();
scene.add(new THREE.HemisphereLight(0xf4f1e8,0x2a302d,0.75));
const key=new THREE.DirectionalLight(0xffffff,0.85);key.position.set(30,60,40);scene.add(key);
const fill=new THREE.DirectionalLight(0xbfd9ff,0.35);fill.position.set(-40,20,-30);scene.add(fill);
const bench=new THREE.Mesh(new THREE.PlaneGeometry(220,160),new THREE.MeshStandardMaterial({color:0xdfe3e0,roughness:.9}));
bench.rotation.x=-Math.PI/2;bench.position.y=-1.6;scene.add(bench);
const grid=new THREE.GridHelper(200,40,0xc3c9c5,0xd2d7d3);grid.position.y=-1.55;scene.add(grid);

const M={
 casing:new THREE.MeshStandardMaterial({color:0xd8c48c,roughness:.65,transparent:true,opacity:1}),
 base:new THREE.MeshStandardMaterial({color:0x4b5052,roughness:.8,transparent:true,opacity:1}),
 lid:new THREE.MeshStandardMaterial({color:0xcbb67c,roughness:.6,transparent:true,opacity:1}),
 cream:new THREE.MeshStandardMaterial({color:0xe6dfca,roughness:.55,transparent:true,opacity:1}),
 black:new THREE.MeshStandardMaterial({color:0x1d2022,roughness:.6}),
 dark:new THREE.MeshStandardMaterial({color:0x34393b,roughness:.5,metalness:.3}),
 alu:new THREE.MeshStandardMaterial({color:0xb7bcc0,roughness:.35,metalness:.7}),
 mirror:new THREE.MeshStandardMaterial({side:THREE.DoubleSide,color:0xeef3f7,roughness:.06,metalness:.95,emissive:0x5d6a74,emissiveIntensity:.55}),
 glass:new THREE.MeshStandardMaterial({color:0xcfe9ff,roughness:.05,metalness:0,transparent:true,opacity:.28,depthWrite:false}),
 quartz:new THREE.MeshStandardMaterial({color:0xe8fbff,roughness:.05,transparent:true,opacity:.35,depthWrite:false}),
 pcb:new THREE.MeshStandardMaterial({color:0x1d6b3e,roughness:.6}),
 chip:new THREE.MeshStandardMaterial({color:0x151515,roughness:.4}),
 copper:new THREE.MeshStandardMaterial({color:0xb87333,roughness:.4,metalness:.8}),
 si:new THREE.MeshStandardMaterial({color:0x2a3d66,roughness:.25,metalness:.6,emissive:0x0b1530}),
 holder:new THREE.MeshStandardMaterial({color:0x202325,roughness:.7}),
};
const root=new THREE.Group();scene.add(root);
const casing=new THREE.Group();root.add(casing);
const comps={};      // cid -> {g, base, exp, info}
const pickables=[];
function reg(cid,g,pos,exp,info){g.position.copy(pos);root.add(g);comps[cid]={g,base:pos.clone(),exp:exp||new THREE.Vector3(),info:info||cid};
  g.traverse(o=>{if(o.isMesh){o.userData.cid=cid;pickables.push(o)}});return g;}
const V=(x,y,z)=>new THREE.Vector3(x,y,z);
function box(w,h,d,m){return new THREE.Mesh(new THREE.BoxGeometry(w,h,d),m)}
function cyl(rt,rb,h,m,seg=24){return new THREE.Mesh(new THREE.CylinderGeometry(rt,rb,h,seg),m)}
function canvasTex(w,h,draw){const c=document.createElement('canvas');c.width=w;c.height=h;const x=c.getContext('2d');draw(x,w,h);const t=new THREE.CanvasTexture(c);t.encoding=THREE.sRGBEncoding;t.anisotropy=4;return {t,c,x};}

/* ---------- exterior (simplified, unbranded) ---------- */
(function exterior(){
 // base tray
 const base=box(55,4,45,M.base);base.position.set(0,2,0);casing.add(base);
 for(const [x,z] of [[-23,18],[23,18],[-23,-18],[23,-18]]){const f=cyl(2,2,1.6,M.base);f.position.set(x,-.8,z);casing.add(f);}
 // body side profile extruded along x (front slopes down)
 const s=new THREE.Shape();
 const P=[[-23.5,4],[-23.5,19.2],[12,20],[23,15.2],[23,4]];
 s.moveTo(-P[0][0],P[0][1]);for(let i=1;i<P.length;i++)s.lineTo(-P[i][0],P[i][1]);s.closePath();
 const geo=new THREE.ExtrudeGeometry(s,{depth:55,bevelEnabled:false});
 const body=new THREE.Mesh(geo,M.casing);body.rotation.y=Math.PI/2;body.position.x=-27.5;casing.add(body);
 // rear vents
 for(let i=0;i<34;i++){const v=box(.5,.25,4,M.base);v.position.set(-15+i*1.25,19.35,-20.5);casing.add(v);}
 // left side fan grille + power switch panel
 for(let i=0;i<7;i++){const v=box(.25,.5,7,M.base);v.position.set(-27.6,7+i*1.5,-17);casing.add(v);}
 const pnl=box(.3,6,6,M.cream);pnl.position.set(-27.65,12,-8.5);casing.add(pnl);
 const sw=box(.6,1.6,2,new THREE.MeshStandardMaterial({color:0xf2f2ef}));
 reg('psw',(()=>{const g=new THREE.Group();g.add(sw);return g})(),V(-27.9,13.5,-8.5),V(0,0,0),'psw');
 casing.add(comps.psw.g);
 // right side lamp-house vent
 for(let i=0;i<8;i++){const v=box(.25,.45,6,M.base);v.position.set(27.6,10+i*1.1,-15);casing.add(v);}
})();
// sample compartment lid (pivot at rear)
const lidPivot=new THREE.Group();lidPivot.position.set(-16,20.1,-21);casing.add(lidPivot);
const lidMesh=box(13,1.4,40,M.lid);lidMesh.position.set(0,.4,20);lidPivot.add(lidMesh);
const lidFront=box(13,9,1.2,M.lid);lidFront.position.set(0,-4,39.6);lidPivot.add(lidFront);
lidPivot.traverse(o=>{if(o.isMesh)o.userData.cid='lid'});
comps.lid={g:lidPivot,base:lidPivot.position.clone(),exp:V(0,0,0),info:'lid',noMove:true};
pickables.push(lidMesh,lidFront);
// console: keypad + display
const lcd=canvasTex(512,320,()=>{});
function drawLCD(lines,bg){const x=lcd.x,w=512,h=320;x.fillStyle=bg||'#b9e6e3';x.fillRect(0,0,w,h);x.fillStyle='#1f3a4a';
 x.font='500 44px IBM Plex Mono, monospace';lines.forEach((l,i)=>x.fillText(l,26,72+i*72));lcd.t.needsUpdate=true;}
drawLCD(['','',' (daya mati)'],'#55605e');
const conG=new THREE.Group();
const disp=new THREE.Group();
const dframe=box(24,15,1.6,M.cream);disp.add(dframe);
const dscr=new THREE.Mesh(new THREE.PlaneGeometry(15,9.4),new THREE.MeshBasicMaterial({map:lcd.t}));dscr.position.z=.81;disp.add(dscr);
disp.position.set(0,8,0);disp.rotation.x=-0.32;conG.add(disp);
const hinge=box(9,3,3,M.cream);hinge.position.set(0,1.2,1);conG.add(hinge);
const kp=canvasTex(512,256,(x,w,h)=>{x.fillStyle='#e7e9e6';x.fillRect(0,0,w,h);x.fillStyle='#9aa3ad';x.fillRect(370,10,130,236);
 const keys=['PRINT','MODE','RETURN','0','1','4','7','.','2','5','8','-','3','6','9','ENTER','GOTO WL','◁','START/STOP','AUTO ZERO','▷'];
 x.font='600 15px IBM Plex Sans, sans-serif';x.textAlign='center';
 let i=0;for(let c=0;c<5;c++)for(let r=0;r<4;r++){if(i>=keys.length)break;const cx=40+c*68,cy=40+r*56;x.fillStyle='#476b8f';x.beginPath();x.ellipse(cx,cy,16,10,0,0,7);x.fill();x.fillStyle='#22313f';x.fillText(keys[i++],cx,cy+28);}
 ['F1','F2','F3','F4'].forEach((k,j)=>{x.fillStyle='#476b8f';x.beginPath();x.ellipse(435,45+j*55,22,13,0,0,7);x.fill();x.fillStyle='#fff';x.fillText(k,435,50+j*55);});});
const kpad=new THREE.Mesh(new THREE.PlaneGeometry(17,8.5),new THREE.MeshStandardMaterial({map:kp.t,roughness:.6}));
kpad.rotation.x=-Math.PI/2+0.42;kpad.position.set(0,.4,8.6);conG.add(kpad);
reg('con',conG,V(11,17.6,9),V(0,22,14),'con');
// I/O ports right side
const ioG=new THREE.Group();
[[0,0,2.2,'#2b2b2b'],[0,0,-1.5,'#2b2b2b'],[0,0,-4,'#a42a2a']].forEach(([x,y,z,c])=>{const b=box(.6,1.1,2.6,new THREE.MeshStandardMaterial({color:c}));b.position.set(x,y,z);ioG.add(b)});
const slot=box(.4,.35,6,M.black);slot.position.set(0,9.8,-4);ioG.add(slot);
reg('io',ioG,V(27.7,2.6,5),V(10,0,0),'io');

/* ---------- optics ---------- */
const Y=8.5; // optical axis height
function lampD2(){const g=new THREE.Group();const hs=box(4.2,6,4.2,M.dark);hs.position.y=-1;g.add(hs);
 const bulb=cyl(1.1,1.1,3.2,M.glass);bulb.position.y=.6;g.add(bulb);
 const win=new THREE.Mesh(new THREE.CircleGeometry(.55,20),new THREE.MeshBasicMaterial({color:0xc9b6ff}));win.position.set(0,.6,1.15);g.add(win);
 const pin=cyl(.15,.15,1.2,M.alu);pin.position.set(.6,-4.5,0);g.add(pin);const pin2=pin.clone();pin2.position.x=-.6;g.add(pin2);
 g.userData.glow=win;return g;}
function lampWI(){const g=new THREE.Group();const sock=box(3,3,3,M.cream);sock.position.y=-2.6;g.add(sock);
 const b=new THREE.Mesh(new THREE.SphereGeometry(1.05,24,16),M.glass);b.scale.set(1,1.35,1);g.add(b);
 const fil=new THREE.Mesh(new THREE.TorusGeometry(.3,.06,6,16),new THREE.MeshBasicMaterial({color:0xffc861}));g.add(fil);
 const sp=box(.2,2.6,2.6,M.alu);sp.position.set(-1.7,-1,0);g.add(sp);g.userData.glow=fil;return g;}
function mirror(w,h,curv){const g=new THREE.Group();const face=new THREE.Group();g.add(face);
 const m=curv?new THREE.Mesh(new THREE.SphereGeometry(8,24,12,-Math.PI/2-w/16,w/8,Math.PI/2-h/16,h/8),M.mirror):new THREE.Mesh(new THREE.PlaneGeometry(w,h),M.mirror);
 if(curv){m.position.z=8.05;}
 face.add(m);const back=box(w+.4,h+.4,.4,M.dark);back.position.z=-.3;face.add(back);
 const post=cyl(.25,.25,Y-1,M.alu);post.position.y=-(Y-1)/2-h/2+.2;g.add(post);g.userData.face=face;return g;}
function slit(){const g=new THREE.Group();const f=new THREE.Group();g.add(f);
 const a=box(1.6,3.2,.15,M.alu);a.position.x=-.88;f.add(a);const b=a.clone();b.position.x=.88;f.add(b);
 const t=box(3.4,.5,.15,M.alu);t.position.y=1.85;f.add(t);const bt=t.clone();bt.position.y=-1.85;f.add(bt);
 const post=box(.4,Y-2,.4,M.dark);post.position.y=-(Y-2)/2-1.9;g.add(post);g.userData.face=f;return g;}
function grating(){const g=new THREE.Group();const face=new THREE.Group();g.add(face);
 const gt=canvasTex(256,256,(x,w,h)=>{const gr=x.createLinearGradient(0,0,w,h);['#6c3bd1','#2c7be5','#1cb37a','#e5c42c','#e2582c','#6c3bd1'].forEach((c,i)=>gr.addColorStop(i/5,c));
  x.globalAlpha=.55;x.fillStyle=gr;x.fillRect(0,0,w,h);x.globalAlpha=1;x.strokeStyle='rgba(20,20,30,.55)';for(let i=0;i<w;i+=3){x.beginPath();x.moveTo(i,0);x.lineTo(i,h);x.stroke();}});
 const surf=new THREE.Mesh(new THREE.SphereGeometry(14,30,10,-Math.PI/2-.17,.34,Math.PI/2-.13,.26),new THREE.MeshStandardMaterial({side:THREE.DoubleSide,map:gt.t,metalness:.85,roughness:.18,emissive:0x333344,emissiveIntensity:.7}));
 surf.position.z=14.05;face.add(surf);
 const back=box(4.9,3.8,.8,M.dark);back.position.z=-.5;face.add(back);
 const drv=cyl(1.4,1.4,2.2,M.alu);drv.position.y=-3.4;g.add(drv);
 const post=cyl(.4,.4,Y-4.5,M.dark);post.position.y=-(Y-4.5)/2-4.4;g.add(post);g.userData.face=face;return g;}
function filterWheel(){const g=new THREE.Group();const f=new THREE.Group();g.add(f);
 const R=2.2;const wheel=new THREE.Group();wheel.position.y=-R;f.add(wheel);
 const disc=cyl(3.4,3.4,.3,M.dark,40);disc.rotation.x=Math.PI/2;wheel.add(disc);
 FILT.forEach((fl,i)=>{const a=Math.PI/2+i*Math.PI/3;const p=new THREE.Mesh(new THREE.CircleGeometry(.75,20),new THREE.MeshStandardMaterial({color:fl[3],transparent:true,opacity:i?0.85:0.15,side:THREE.DoubleSide,emissive:fl[3],emissiveIntensity:.25}));
  p.position.set(Math.cos(a)*R,Math.sin(a)*R,.17);wheel.add(p);});
 const hub=cyl(.6,.6,1.4,M.alu);hub.rotation.x=Math.PI/2;wheel.add(hub);
 const mot=cyl(.9,.9,2,M.alu);mot.rotation.x=Math.PI/2;mot.position.set(0,-R,-1.4);f.add(mot);
 g.userData.face=f;g.userData.wheel=wheel;return g;}
function splitter(){const g=new THREE.Group();const face=new THREE.Group();g.add(face);
 const p=box(3.2,3.2,.18,new THREE.MeshStandardMaterial({color:0xbfe8ff,metalness:.6,roughness:.08,transparent:true,opacity:.55,emissive:0x2b4a60}));face.add(p);
 const fr=box(3.6,.35,.5,M.dark);fr.position.y=1.75;face.add(fr);const fr2=fr.clone();fr2.position.y=-1.75;face.add(fr2);
 const post=cyl(.25,.25,Y-2,M.alu);post.position.y=-(Y-2)/2-1.8;g.add(post);g.userData.face=face;return g;}
function windowPane(){const g=new THREE.Group();const f=new THREE.Group();g.add(f);
 const p=box(2.6,2.6,.2,M.quartz);f.add(p);const fr=new THREE.Mesh(new THREE.TorusGeometry(1.75,.18,6,4),M.dark);fr.rotation.z=Math.PI/4;f.add(fr);g.userData.face=f;return g;}
function lens(){const g=new THREE.Group();const f=new THREE.Group();g.add(f);
 const l=new THREE.Mesh(new THREE.SphereGeometry(1.2,24,16),M.quartz);l.scale.set(1,1,.32);f.add(l);
 const r=new THREE.Mesh(new THREE.TorusGeometry(1.25,.15,8,24),M.dark);f.add(r);
 const post=box(.3,Y-3,.3,M.dark);post.position.y=-(Y-3)/2-1.3;g.add(post);g.userData.face=f;return g;}
function photodiode(){const g=new THREE.Group();const f=new THREE.Group();g.add(f);
 const can=cyl(1,1,1.1,M.alu);can.rotation.x=Math.PI/2;f.add(can);
 const win=new THREE.Mesh(new THREE.PlaneGeometry(1.1,1.1),M.si);win.position.z=.57;f.add(win);
 for(const dx of[-.35,.35]){const lg=cyl(.06,.06,1.6,M.copper,6);lg.rotation.x=Math.PI/2;lg.position.set(dx,0,-1.3);f.add(lg);}
 return g;}
function cellHolder(liquidCol){const g=new THREE.Group();
 const base=box(3.6,1,3.6,M.holder);base.position.y=-3.2;g.add(base);
 const wl=box(.5,4.6,3.6,M.holder);wl.position.set(0,-.6,1.55);// front & back walls (z)
 const w1=box(3.6,4.6,.5,M.holder);w1.position.set(0,-.6,1.55);g.add(w1);const w2=w1.clone();w2.position.z=-1.55;g.add(w2);
 const knob=cyl(.45,.45,1,M.alu);knob.position.set(0,2.2,2.4);g.add(knob);
 const cuv=new THREE.Group();cuv.position.y=.4;
 const shell=box(1.25,4.5,1.25,M.quartz);cuv.add(shell);
 const liq=box(1.05,3.2,1.05,new THREE.MeshStandardMaterial({color:liquidCol||0xffffff,transparent:true,opacity:.5,depthWrite:false}));liq.position.y=-.55;cuv.add(liq);
 g.add(cuv);g.userData.cuv=cuv;g.userData.liq=liq;return g;}
function pcb(w,d,parts){const g=new THREE.Group();const b=box(w,.18,d,M.pcb);g.add(b);
 (parts||[]).forEach(([x,z,sw,sd,h,m])=>{const c=box(sw,h,sd,m||M.chip);c.position.set(x,h/2+.09,z);g.add(c)});return g;}

const XS=-17; // cell holder x
reg('d2',lampD2(),V(20,Y,-19.5),V(6,12,-6),'d2');
reg('wi',lampWI(),V(24.5,Y,-13.5),V(10,12,0),'wi');
reg('m1',mirror(2.4,2.4,true),V(20,Y,-13.5),V(4,10,0),'m1');
reg('s1',slit(),V(14,Y,-13.5),V(0,8,0),'s1');
reg('g',grating(),V(3,Y,-19),V(0,10,-6),'g');
reg('s2',slit(),V(9.5,Y,-8.8),V(0,8,2),'s2');
reg('f',filterWheel(),V(10.6,Y,-7.4),V(3,8,3),'f');
reg('m2',mirror(2.8,2.6,true),V(13.5,Y,-4),V(4,8,4),'m2');
reg('m3',splitter(),V(-1.5,Y,-2.2),V(0,8,0),'m3');
reg('m4',mirror(2.4,2.4,false),V(-8,Y,3),V(0,6,6),'m4');
reg('m5',mirror(2.4,2.4,false),V(-8,Y,-6.5),V(0,6,-6),'m5');
reg('wS',windowPane(),V(-10.5,Y,3),V(0,4,6),'w');
reg('wR',windowPane(),V(-10.5,Y,-6.5),V(0,4,-6),'w');
reg('cellS',cellHolder(),V(XS,Y,3),V(0,9,6),'cellS');
reg('cellR',cellHolder(),V(XS,Y,-6.5),V(0,9,-6),'cellR');
reg('lS',lens(),V(-23.4,Y,3),V(-3,4,6),'l');
reg('lR',lens(),V(-23.4,Y,-6.5),V(-3,4,-6),'l');
reg('pdS',photodiode(),V(-25.6,Y,3),V(-7,4,6),'pd');
reg('pdR',photodiode(),V(-25.6,Y,-6.5),V(-7,4,-6),'pd');
// compartment floor & walls (inside casing, conceptual partition)
const comp=new THREE.Group();const fl=box(13,.3,36,M.holder);fl.position.set(-16,4.4,-1);comp.add(fl);
const wR_=box(.3,14,36,new THREE.MeshStandardMaterial({color:0x2a2e30,transparent:true,opacity:.35}));wR_.position.set(-9.6,11,-1);comp.add(wR_);
const wL_=wR_.clone();wL_.position.x=-22.4;comp.add(wL_);casing.add(comp);
// electronics
const pre=pcb(8,.1,[]);pre.children[0].scale.set(1,1,1);
const preG=new THREE.Group();const preB=box(.18,6,13,M.pcb);preG.add(preB);
[[0,1,2],[0,-1.5,-3],[0,1.5,-5]].forEach(([x,y,z])=>{const c=box(.6,1,1.2,M.chip);c.position.set(.35,y,z);preG.add(c)});
reg('pre',preG,V(-26.9,9,-1.8),V(-10,0,0),'pre');
reg('cpu',pcb(14,16,[[-3,-3,3,3,.5],[2,2,4,2,.4],[4,-5,1.5,1.5,.9],[-5,5,2,1,.4],[0,6,6,.8,.6,M.alu]]),V(19,5.2,8),V(14,-4,6),'cpu');
const pwrG=new THREE.Group();const tr=box(5,4.5,5,M.dark);tr.position.y=2.3;pwrG.add(tr);const coil=cyl(2,2,3,M.copper);coil.rotation.z=Math.PI/2;coil.position.set(0,2.4,0);coil.scale.set(.9,.9,.9);pwrG.add(coil);
const pb=pcb(8,9,[[1,1,2,2,.8],[-2,-2,1.2,1.2,1.4,M.alu]]);pb.position.set(9,0,0);pwrG.add(pb);
reg('pwr',pwrG,V(-3,4.2,-15),V(0,-4,-14),'pwr');
const fanG=new THREE.Group();const fr=box(1.2,8,8,M.black);fanG.add(fr);const bl=cyl(3.3,3.3,.4,M.dark,7);bl.rotation.z=Math.PI/2;bl.position.x=.5;fanG.add(bl);
reg('fan',fanG,V(-26.6,11.5,-17),V(-12,0,-4),'fan');

/* ---------- beams ---------- */
const beamGroup=new THREE.Group();scene.add(beamGroup);
const segGeo=new THREE.CylinderGeometry(.13,.13,1,8,1,true);
const glowTex=canvasTex(64,64,(x)=>{const g=x.createRadialGradient(32,32,0,32,32,32);g.addColorStop(0,'rgba(255,255,255,1)');g.addColorStop(.35,'rgba(255,255,255,.45)');g.addColorStop(1,'rgba(255,255,255,0)');x.fillStyle=g;x.fillRect(0,0,64,64)}).t;
const segs=[];
function addSeg(a,b,stage,kind){const mat=new THREE.MeshBasicMaterial({color:0xffffff,transparent:true,opacity:.85,depthWrite:false,blending:THREE.AdditiveBlending});
 const m=new THREE.Mesh(segGeo,mat);beamGroup.add(m);
 const ph=[];for(let i=0;i<2;i++){const s=new THREE.Sprite(new THREE.SpriteMaterial({map:glowTex,color:0xffffff,transparent:true,depthWrite:false,blending:THREE.AdditiveBlending}));s.scale.set(1.3,1.3,1);beamGroup.add(s);ph.push(s);}
 segs.push({a,b,stage,kind,m,ph});}
addSeg('LAMP','m1',1,'poly');addSeg('m1','s1',2,'poly');addSeg('s1','g',3,'poly');
addSeg('g','s2',4,'mono');addSeg('s2','f',5,'mono');addSeg('f','m2',6,'mono');addSeg('m2','m3',7,'mono');
addSeg('m3','m4',8,'S');addSeg('m4','wS',9,'S');addSeg('wS','cellS',10,'S');addSeg('cellS','lS',11,'S');addSeg('lS','pdS',12,'S');
addSeg('m3','m5',8,'R');addSeg('m5','wR',9,'R');addSeg('wR','cellR',10,'R');addSeg('cellR','lR',11,'R');addSeg('lR','pdR',12,'R');
// dispersion fan at grating
const fan=[];for(let i=0;i<9;i++){const m=new THREE.Mesh(segGeo,new THREE.MeshBasicMaterial({transparent:true,opacity:.35,depthWrite:false,blending:THREE.AdditiveBlending}));beamGroup.add(m);fan.push(m);}

function wl2rgb(l){let r=0,g=0,b=0;
 if(l<380){return new THREE.Color(0.62,0.42,1)}
 if(l<440){r=-(l-440)/60;b=1}else if(l<490){g=(l-440)/50;b=1}else if(l<510){g=1;b=-(l-510)/20}else if(l<580){r=(l-510)/70;g=1}else if(l<645){r=1;g=-(l-645)/65}else if(l<=780){r=1}else{return new THREE.Color(0.75,0.12,0.12)}
 let f=1;if(l<420)f=.3+.7*(l-380)/40;else if(l>700)f=.3+.7*(780-l)/80;return new THREE.Color(r*f+.05,g*f+.05,b*f+.05);}

/* ---------- state ---------- */
const S={wl:540,sw:340.8,explode:0,expT:0,caseMode:'cut',labels:innerWidth<700?'code':'full',beams:!matchMedia('(prefers-reduced-motion: reduce)').matches||true,
 step:false,stepI:0,sel:null,lidOpen:0,lidT:0};
const tmp=new THREE.Vector3(),tmp2=new THREE.Vector3();
function anchor(id,out){if(id==='LAMP')id=activeLamp();return comps[id].g.getWorldPosition(out||new THREE.Vector3());}
function activeLamp(){if(S.forceLamp)return S.forceLamp;return S.wl<S.sw?'d2':'wi'}
function filterIdx(l){for(let i=0;i<FILT.length;i++)if(l<=FILT[i][2]+.99)return i;return 5}
function orientBisect(id,prevId,nextId,extraY){const c=comps[id];const p=anchor(id),a=anchor(prevId),b=anchor(nextId);
 const n=a.sub(p).normalize().add(b.sub(p).normalize()).normalize();const face=c.g.userData.face||c.g;
 const t=new THREE.Object3D();t.position.copy(p);t.lookAt(p.clone().add(n));
 const q=t.quaternion.clone();if(extraY){q.multiply(new THREE.Quaternion().setFromAxisAngle(V(0,1,0),extraY));}
 return q;}
function faceToward(id,targetId){const c=comps[id];const p=anchor(id),b=anchor(targetId);const t=new THREE.Object3D();t.position.copy(p);t.lookAt(b);return t.quaternion;}
function applyQ(id,q,lerp){const c=comps[id];const face=c.g.userData.face||c.g;
 // face is child of group (group has no rotation) -> world quat == local
 if(lerp)face.quaternion.slerp(q,lerp);else face.quaternion.copy(q);}

/* ---------- labels ---------- */
const lblBox=document.getElementById('labels'),svg=document.getElementById('leaders');
const LBL=['d2','wi','m1','s1','g','s2','f','m2','m3','m4','m5','wS','wR','cellS','cellR','lS','lR','pdS','pdR','pre','cpu','pwr','fan'];
const EXT=['lid','con','io','psw'];
const labels={};
[...LBL,...EXT].forEach(cid=>{const inf=INFO[comps[cid].info];const d=document.createElement('div');d.className='lbl';d.dataset.cid=cid;
 const suffix=cid.endsWith('S')&&cid.length>1&&cid!=='cellS'?' (S)':cid.endsWith('R')&&cid.length>1&&cid!=='cellR'?' (R)':'';
 d.innerHTML=`<b><span class="c">${inf.code}</span>${inf.id}${suffix}</b><i>${inf.en}</i>`;d.dataset.code=(inf.code||inf.id.split(' ')[0])+suffix;
 d.addEventListener('click',e=>{e.stopPropagation();select(cid,true)});lblBox.appendChild(d);
 const ln=document.createElementNS('http://www.w3.org/2000/svg','line');ln.setAttribute('stroke','rgba(200,230,225,.7)');ln.setAttribute('stroke-width','1');svg.appendChild(ln);
 labels[cid]={d,ln,ext:EXT.includes(cid)};});
function setLabelMode(){Object.values(labels).forEach(L=>{const inf=L.d;if(S.labels==='code'){inf.innerHTML=`<b>${inf.dataset.code}</b>`;}else if(S.labels==='full'){const c=comps[inf.dataset.cid];const I=INFO[c.info];const sfx=inf.dataset.code.match(/\((S|R)\)$/)?' ('+RegExp.$1+')':'';inf.innerHTML=`<b><span class="c">${I.code}</span>${I.id}${sfx}</b><i>${I.en}</i>`;}});
 document.getElementById('bLbl').textContent='Label: '+({full:'lengkap',code:'kode',off:'mati'})[S.labels];}

/* ---------- picking & selection ---------- */
const ray=new THREE.Raycaster(),mouse=new THREE.Vector2();let downPos=null;
renderer.domElement.addEventListener('pointerdown',e=>{downPos=[e.clientX,e.clientY]});
renderer.domElement.addEventListener('pointerup',e=>{if(!downPos)return;const d=Math.hypot(e.clientX-downPos[0],e.clientY-downPos[1]);downPos=null;if(d>6)return;
 const r=renderer.domElement.getBoundingClientRect();mouse.set((e.clientX-r.left)/r.width*2-1,-(e.clientY-r.top)/r.height*2+1);ray.setFromCamera(mouse,camera);
 const vis=pickables.filter(o=>{let p=o;while(p){if(!p.visible)return false;p=p.parent}return true});
 const hits=ray.intersectObjects(vis,false).filter(h=>!(S.caseMode!=='solid'&&['lid'].includes(h.object.userData.cid)&&false));
 if(hits.length)select(hits[0].object.userData.cid,false);});
const selBox=new THREE.BoxHelper(undefined,0x4cc2c4);selBox.visible=false;scene.add(selBox);
function select(cid,focus){S.sel=cid;const c=comps[cid];selBox.setFromObject(c.g);selBox.visible=true;
 Object.values(labels).forEach(L=>L.d.classList.toggle('sel',L.d.dataset.cid===cid));renderInfo(cid);
 if(focus){const p=anchor(cid);flyTo(p);}}
let fly=null;
function flyTo(p,dist,dir){dir=dir?dir.clone().normalize():camera.position.clone().sub(controls.target).normalize();dist=dist||Math.min(camera.position.distanceTo(controls.target),48);
 fly={t0:performance.now(),fromT:controls.target.clone(),toT:p.clone(),fromP:camera.position.clone(),toP:p.clone().add(dir.multiplyScalar(dist))};}
function renderInfo(cid){const box=document.getElementById('info');if(!cid){box.innerHTML=`<div class="eyebrow">Komponen terpilih</div><h2 style="margin-top:4px">Ketuk komponen pada model</h2><p>Model menampilkan seluruh rantai optik double-beam dari lampu hingga fotodioda. Geser λ untuk melihat pergantian lampu, posisi filter, dan rotasi kisi; aktifkan mode bertahap untuk menelusuri jalur cahaya satu per satu.</p>`;return}
 const I=INFO[comps[cid].info];
 box.innerHTML=`<div class="eyebrow">Komponen terpilih</div><h2 style="margin-top:4px"><span class="code">${I.code}</span> ${I.id}</h2><div style="color:var(--muted);margin-top:-2px">${I.en}</div>
 <p style="margin-top:8px"><b>Fungsi.</b> ${I.fn}</p><p><b>Prinsip kerja.</b> ${I.pr}</p>
 <p style="font-size:12.5px"><span class="pill ${I.st}">${I.st==='v'?'Terverifikasi':'Konseptual'}</span><span class="mono" style="color:var(--muted)">${I.src}</span></p>`;}

/* ---------- step mode ---------- */
const STEPS=[
 {c:'LAMP',t:'Sumber cahaya',d:()=>activeLamp()==='d2'?`λ ${S.wl.toFixed(1)} nm < ${S.sw} nm: lampu D₂ aktif, memancarkan kontinum UV.`:`λ ${S.wl.toFixed(1)} nm ≥ ${S.sw} nm: lampu halogen WI aktif.`},
 {c:'m1',t:'M1',d:()=>'Cermin pemilih sumber berputar ke lampu aktif dan memantulkan cahaya polikromatik ke celah masuk.'},
 {c:'s1',t:'S1',d:()=>'Celah masuk membentuk berkas sempit yang masuk ke monokromator.'},
 {c:'g',t:'G',d:()=>'Kisi mendispersi cahaya; rotasi kisi memilih λ yang terfokus di celah keluar. Kipas warna menunjukkan dispersi secara konseptual.'},
 {c:'s2',t:'S2',d:()=>`Celah keluar meloloskan pita 1 nm di sekitar ${S.wl.toFixed(1)} nm.`},
 {c:'f',t:'F',d:()=>`Filter ${FILT[filterIdx(S.wl)][0]} berada di jalur untuk menahan orde dua/tiga.`},
 {c:'m2',t:'M2',d:()=>'Cermin toroidal mengarahkan berkas monokromatik ke pembagi berkas.'},
 {c:'m3',t:'M3',d:()=>'Half mirror membagi berkas menjadi berkas sampel dan referensi.'},
 {c:'m4',t:'M4/M5',d:()=>'Cermin sisi S dan R membelokkan masing-masing berkas ke kompartemen sampel.'},
 {c:'wS',t:'W',d:()=>'Berkas melewati jendela kuarsa dan terfokus di pusat pemegang kuvet.'},
 {c:'cellS',t:'Kuvet S/R',d:()=>'Analit di kuvet S menyerap sebagian radiasi (I); kuvet R berisi blanko (I₀).'},
 {c:'lS',t:'L',d:()=>'Lensa mengumpulkan berkas ke fotodioda.'},
 {c:'pdS',t:'P.D.',d:()=>'Dua fotodioda Si mengubah intensitas S dan R menjadi arus foto.'},
 {c:'pre',t:'Preamp',d:()=>'Preamplifier mengubah arus menjadi tegangan.'},
 {c:'cpu',t:'CPU',d:()=>'CPU mengonversi A/D dan menghitung A ≈ log₁₀(I_R/I_S) setelah koreksi baseline.'}];
const cap=document.getElementById('caption');
function showStep(){const s=STEPS[S.stepI];cap.style.display=S.step?'block':'none';if(!S.step)return;
 cap.innerHTML=`<b>Tahap ${S.stepI+1}/${STEPS.length} · ${s.t}</b> — ${s.d()}`;const cid=s.c==='LAMP'?activeLamp():s.c;select(cid,false);
 document.querySelectorAll('#flow button').forEach((b,i)=>b.classList.toggle('on',i===S.stepI));
 document.getElementById('flowtxt').textContent=`${s.t}: ${s.d()}`;}
const flow=document.getElementById('flow');
STEPS.forEach((s,i)=>{if(i)flow.insertAdjacentHTML('beforeend','<span class="ar">→</span>');const b=document.createElement('button');b.textContent=s.t;b.onclick=()=>{S.step=true;S.stepI=i;syncBtns();showStep();const cid=s.c==='LAMP'?activeLamp():s.c;flyTo(anchor(cid));};flow.appendChild(b);});
document.getElementById('ftab').innerHTML=FILT.map(f=>`<tr><td>${f[0]}</td><td class="n">${f[1]}–${f[2]}</td></tr>`).join('');

/* ---------- toolbar ---------- */
const $=id=>document.getElementById(id);
function syncBtns(){$('bCase').setAttribute('aria-pressed',S.caseMode==='solid');$('bCut').setAttribute('aria-pressed',S.caseMode==='cut');
 $('bExp').setAttribute('aria-pressed',S.explode===1);$('bBeam').setAttribute('aria-pressed',S.beams);$('bStep').setAttribute('aria-pressed',S.step);
 $('bPrev').hidden=$('bNext').hidden=!S.step;
 $('bCase').textContent=S.caseMode==='hidden'?'Casing disembunyikan':'Casing tampak';}
$('bCase').onclick=()=>{S.caseMode=S.caseMode==='solid'?'hidden':'solid';applyCase();syncBtns();};
$('bCut').onclick=()=>{S.caseMode=S.caseMode==='cut'?'solid':'cut';applyCase();syncBtns();};
$('bExp').onclick=()=>{S.explode=S.explode?0:1;syncBtns();};
$('bLbl').onclick=()=>{S.labels={full:'code',code:'off',off:'full'}[S.labels];setLabelMode();};
$('bBeam').onclick=()=>{S.beams=!S.beams;syncBtns();};
$('bStep').onclick=()=>{S.step=!S.step;syncBtns();showStep();};
$('bPrev').onclick=()=>{S.stepI=(S.stepI+STEPS.length-1)%STEPS.length;showStep();};
$('bNext').onclick=()=>{S.stepI=(S.stepI+1)%STEPS.length;showStep();};
$('bReset').onclick=()=>{fly=null;resetCam();};
function applyCase(){const op=S.caseMode==='cut'?0.16:1;[M.casing,M.base,M.lid,M.cream].forEach(m=>{m.opacity=op;m.depthWrite=op===1;m.needsUpdate=true});
 casing.visible=S.caseMode!=='hidden';comps.con.g.visible=S.caseMode!=='hidden';comps.io.g.visible=S.caseMode!=='hidden';
 [M.casing,M.lid,M.cream].forEach(m=>m.transparent=true);}
$('wl').oninput=e=>{S.wl=+e.target.value;$('wlv').textContent=S.wl.toFixed(1)+' nm';updateStat();if(S.step)showStep();};
$('sw').onchange=e=>{let v=Math.min(364,Math.max(295,+e.target.value||340.8));e.target.value=v;S.sw=v;updateStat();};
function updateStat(){$('stat').textContent=`Lampu: ${activeLamp()==='d2'?'D₂':'WI'} · Filter: ${FILT[filterIdx(S.wl)][0]}`;}

/* ---------- loop ---------- */
function resize(){const w=stage.clientWidth,h=stage.clientHeight;renderer.setSize(w,h,false);camera.aspect=w/h;camera.updateProjectionMatrix();}
new ResizeObserver(resize).observe(stage);resize();
const casingLift=V(0,0,0);let last=performance.now();
const reduce=matchMedia('(prefers-reduced-motion: reduce)').matches;
function frame(now){const dt=Math.min(.05,(now-last)/1000);last=now;simTick(dt);
 if(fly){const k=Math.min(1,(now-fly.t0)/700),e=k<.5?2*k*k:1-Math.pow(-2*k+2,2)/2;controls.target.lerpVectors(fly.fromT,fly.toT,e);camera.position.lerpVectors(fly.fromP,fly.toP,e);if(k>=1)fly=null;}
 controls.update();
 S.expT+=(S.explode-S.expT)*Math.min(1,dt*5);
 for(const id in comps){const c=comps[id];if(c.noMove)continue;c.g.position.copy(c.base).addScaledVector(c.exp,S.expT);}
 casing.position.y=S.expT*28;
 S.lidT+=(S.lidOpen-S.lidT)*Math.min(1,dt*5);lidPivot.rotation.x=-S.lidT*1.25;
 root.updateMatrixWorld(true);casing.updateMatrixWorld(true);
 // optics orientation
 const lamp=activeLamp();
 applyQ('m1',orientBisect('m1',lamp,'s1'),reduce?1:Math.min(1,dt*4));
 const gx=(S.wl-550)/910*0.6; // conceptual rotation
 applyQ('g',orientBisect('g','s1','s2',gx),reduce?1:Math.min(1,dt*6));
 applyQ('m2',orientBisect('m2','f','m3'));applyQ('m3',orientBisect('m3','m2','m5'));
 applyQ('m4',orientBisect('m4','m3','wS'));applyQ('m5',orientBisect('m5','m3','wR'));
 ['s1','s2','f'].forEach(id=>applyQ(id,faceToward(id,{s1:'g',s2:'f',f:'m2'}[id])));
 ['wS','lS'].forEach(id=>applyQ(id,faceToward(id,'pdS')));['wR','lR'].forEach(id=>applyQ(id,faceToward(id,'pdR')));
 applyQ('pdS',faceToward('pdS','lS'));applyQ('pdR',faceToward('pdR','lR'));
 const wh=comps.f.g.userData.wheel;const fi=filterIdx(S.wl);const target=-fi*Math.PI/3;wh.rotation.z+=(target-wh.rotation.z)*Math.min(1,dt*(reduce?50:5));
 const lon=S.lampsOn!==false;comps.d2.g.userData.glow.material.color.set(lon&&lamp==='d2'?0xd8c8ff:0x333040);comps.wi.g.userData.glow.material.color.set(lon&&lamp==='wi'?0xffd06b:0x403a2a);
 // beams
 const mono=wl2rgb(S.wl),poly=new THREE.Color(lamp==='d2'?0x8b5cf6:0xe0a21b),cS=new THREE.Color(0x13c4b5),cR=new THREE.Color(0xff7a45);
 const tt=now/1000;
 segs.forEach((s,i)=>{const show=S.beams&&(!S.step||s.stage<=S.stepI+0)&&!(S.step&&S.stepI===0&&false);
  const vis=S.beams&&(!S.step||s.stage<=Math.max(1,S.stepI));
  s.m.visible=vis;s.ph.forEach(p=>p.visible=vis&&!reduce);if(!vis)return;
  const a=anchor(s.a,tmp),b=anchor(s.b,tmp2);const d=b.clone().sub(a);const L=d.length();
  s.m.position.copy(a).addScaledVector(d,.5);s.m.scale.set(1,L,1);s.m.quaternion.setFromUnitVectors(V(0,1,0),d.normalize());
  const col=s.kind==='poly'?poly:s.kind==='mono'?mono:s.kind==='S'?cS:cR;s.m.material.color.copy(col);
  // attenuation in sample cell (conceptual)
  if(s.kind==='S'&&s.stage>=11)s.m.material.opacity=.55;else s.m.material.opacity=.85;
  s.ph.forEach((p,k)=>{const u=((tt*0.9+k*.5+i*.13)%1);p.position.copy(a).addScaledVector(b.clone().sub(a),u);p.material.color.copy(col);});});
 const fanVis=S.beams&&(!S.step||S.stepI>=3);const gp=anchor('g'),s2p=anchor('s2'),f2=anchor('f');
 const axis=s2p.clone().sub(gp);const perp=V(axis.z,0,-axis.x).normalize();
 fan.forEach((m,i)=>{m.visible=fanVis;if(!fanVis)return;const lam=390+i*45;const off=(lam-S.wl)/110*1.6;const end=s2p.clone().addScaledVector(perp,off);
  const d=end.clone().sub(gp);m.position.copy(gp).addScaledVector(d,.5);m.scale.set(.6,d.length(),.6);m.quaternion.setFromUnitVectors(V(0,1,0),d.clone().normalize());m.material.color.copy(wl2rgb(lam));m.material.opacity=Math.abs(off)<.3?0:.28;});
 if(S.sel&&selBox.visible)selBox.setFromObject(comps[S.sel].g);
 renderer.render(scene,camera);updateLabels();requestAnimationFrame(frame);}
const ctr=new THREE.Vector3(0,8,-2);
function updateLabels(){const w=stage.clientWidth,h=stage.clientHeight;const c=ctr.clone().project(camera);const cx=(c.x+1)/2*w,cy=(1-c.y)/2*h;
 const inside=S.caseMode!=='solid'||S.expT>.5;const placed=[];
 for(const cid in labels){const L=labels[cid];let show=S.labels!=='off'&&(L.ext?S.caseMode!=='hidden':inside)&&(cid==='lid'?S.caseMode!=='hidden':true);
  if(S.step&&!L.ext){const st=STEPS[S.stepI];const cur=st.c==='LAMP'?activeLamp():st.c;const pairs={m4:['m4','m5'],wS:['wS','wR'],cellS:['cellS','cellR'],lS:['lS','lR'],pdS:['pdS','pdR']};show=show&&((pairs[cur]||[cur]).includes(cid));}
  if(!show){L.d.style.display='none';L.ln.style.display='none';continue}
  const p=anchor(cid,tmp);if(cid==='lid'){p.set(-16,22+S.expT*28,8)}if(cid==='con'){p.y+=6}
  const v=p.clone().project(camera);if(v.z>1){L.d.style.display='none';L.ln.style.display='none';continue}
  const ax=(v.x+1)/2*w,ay=(1-v.y)/2*h;let dx=ax-cx,dy=ay-cy;const dl=Math.hypot(dx,dy)||1;dx/=dl;dy/=dl;
  const len=S.labels==='code'?34:58;let lx=ax+dx*len,ly=ay+dy*len-(L.ext?0:0);
  // simple collision push
  const bw=S.labels==='code'?44:150,bh=S.labels==='code'?18:30;
  for(let k=0;k<8;k++){const hit=placed.find(q=>Math.abs(q[0]-lx)<(bw+q[2])/2&&Math.abs(q[1]-ly)<(bh+q[3])/2);if(!hit)break;ly+=(ly>=hit[1]?1:-1)*(bh*.6);}
  lx=Math.max(bw/2+4,Math.min(w-bw/2-4,lx));ly=Math.max(bh/2+22,Math.min(h-bh/2-4,ly));placed.push([lx,ly,bw,bh]);
  L.d.style.display='block';L.d.style.left=lx+'px';L.d.style.top=ly+'px';
  L.ln.style.display='block';L.ln.setAttribute('x1',ax);L.ln.setAttribute('y1',ay);L.ln.setAttribute('x2',lx);L.ln.setAttribute('y2',ly);}
}
applyCase();setLabelMode();syncBtns();updateStat();renderInfo(null);requestAnimationFrame(frame);

/* ---------- component table ---------- */
const order=[['d2','d2'],['wi','wi'],['m1','m1'],['s1','s1'],['g','g'],['s2','s2'],['f','f'],['m2','m2'],['m3','m3'],['m4','m4'],['m5','m5'],['wS','w'],['cellS','cellS'],['cellR','cellR'],['lS','l'],['pdS','pd'],['pre','pre'],['cpu','cpu'],['pwr','pwr'],['fan','fan'],['con','con'],['lid','lid'],['io','io'],['psw','psw']];
document.querySelector('#ktab tbody').innerHTML=order.map(([cid,k])=>{const I=INFO[k];return `<tr class="clk" data-cid="${cid}"><td class="mono">${I.code||'—'}</td><td><b>${I.id}</b><br><span style="color:var(--muted)">${I.en}</span></td><td>${I.fn}</td><td><span class="pill ${I.st}">${I.st==='v'?'Terverifikasi':'Konseptual'}</span></td><td class="mono" style="font-size:11.5px">${I.src}</td></tr>`}).join('');
document.querySelectorAll('#ktab tr.clk').forEach(r=>r.onclick=()=>{const cid=r.dataset.cid;if(['pre','cpu','pwr','fan'].includes(cid)&&S.caseMode==='solid'){S.caseMode='cut';applyCase();syncBtns();}select(cid,true);stage.scrollIntoView({behavior:'smooth',block:'nearest'});});

/* ---------- tabs ---------- */
document.querySelectorAll('.tab').forEach(t=>t.onclick=()=>{document.querySelectorAll('.tab').forEach(x=>x.setAttribute('aria-selected',x===t));document.querySelectorAll('[role=tabpanel]').forEach(p=>p.hidden=p.id!==t.dataset.t);});

/* ---------- practicum simulation (IK/LAB.INST/6) — pemutaran otomatis ---------- */
const STD=[[0,0.000],[0.1,0.080],[0.2,0.155],[0.3,0.234],[0.4,0.312],[0.5,0.380]];
const SOLS=[{k:'blk',n:'Blanko (pelarut/pereaksi)',a:0,col:0xf4f4f4},...STD.map(([c,a],i)=>({k:'std'+i,n:`Standar ${c.toFixed(1)} mg/L`,a,c,col:0xd04fb0,int:a/0.38})),
 {k:'bs',n:'Blanko sampel',a:-0.001,col:0xf2eef2},{k:'u1',n:'Sampel 1',a:0.191,col:0xd04fb0,int:.5},{k:'u2',n:'Sampel 2',a:0.184,col:0xd04fb0,int:.48},{k:'u3',n:'Sampel 3 + spike',a:0.193,col:0xd04fb0,int:.51}];
const solBy=k=>SOLS.find(s=>s.k===k);
const P={S:null,R:null,power:false,read:[],unk:[]};
// kantong silika gel di kompartemen
const silica=new THREE.Mesh(new THREE.BoxGeometry(5,1.2,3.5),new THREE.MeshStandardMaterial({color:0xf3f1ea,roughness:.9}));silica.position.set(-16,5.2,12);scene.add(silica);
function setCuv(slot,k){P[slot]=k;const cell=comps[slot==='S'?'cellS':'cellR'].g;const s=solBy(k);const liq=cell.userData.liq;
 cell.userData.cuv.visible=!!s;liq.visible=!!s;if(s){const c=new THREE.Color(s.col);if(s.int!=null){c.lerp(new THREE.Color(0xffffff),1-Math.min(1,s.int*1.2));}liq.material.color.copy(c);liq.material.opacity=s.int?0.35+0.4*s.int:0.25;}
 $('slots').textContent=`S (depan): ${P.S?solBy(P.S).n:'kosong'} · R (belakang): ${P.R?solBy(P.R).n:'kosong'}`;}
function log(t){$('log').textContent+=t+'\n';$('log').scrollTop=1e6;}
let WL=540,SCAN=[700,400];
function regress(){const pts=P.read;const n=pts.length;if(n<2)return null;const mx=pts.reduce((s,p)=>s+p[0],0)/n,my=pts.reduce((s,p)=>s+p[1],0)/n;
 let sxx=0,sxy=0,syy=0;pts.forEach(([x,y])=>{sxx+=(x-mx)**2;sxy+=(x-mx)*(y-my);syy+=(y-my)**2});const b=sxy/sxx,a=my-b*mx;return {a,b,r:sxy/Math.sqrt(sxx*syy)};}
function drawCal(){const g=$('cal');const W=360,H=170,l=44,r=12,t=10,bm=30;const xs=v=>l+v/0.5*(W-l-r),ys=v=>H-bm-v/0.4*(H-t-bm);
 let s='';
 for(let v=0;v<=0.4001;v+=0.1){s+=`<line x1="${l}" x2="${W-r}" y1="${ys(v)}" y2="${ys(v)}" stroke="var(--line)"/><text x="${l-6}" y="${ys(v)+4}" font-size="10" text-anchor="end" fill="var(--muted)">${v.toFixed(1)}</text>`;}
 for(let v=0;v<=0.5001;v+=0.1)s+=`<text x="${xs(v)}" y="${H-bm+14}" font-size="10" text-anchor="middle" fill="var(--muted)">${v.toFixed(1)}</text>`;
 s+=`<text x="${(l+W-r)/2}" y="${H-3}" font-size="10" text-anchor="middle" fill="var(--muted)">Konsentrasi (mg/L)</text><text x="10" y="${t+70}" font-size="10" fill="var(--muted)" transform="rotate(-90 10 ${t+70})" text-anchor="middle">Absorbansi</text>`;
 const R=regress();if(R){s+=`<line x1="${xs(0)}" y1="${ys(R.a)}" x2="${xs(0.5)}" y2="${ys(R.a+R.b*0.5)}" stroke="var(--accent)" stroke-width="1.5"/>`;
  s+=`<text x="${l+6}" y="${t+12}" font-size="10.5" fill="var(--ink)">A = ${R.b.toFixed(4)}·C ${R.a<0?'−':'+'} ${Math.abs(R.a).toFixed(4)}   r = ${R.r.toFixed(4)}</text>`;}
 P.read.forEach(([x,y])=>s+=`<circle cx="${xs(x)}" cy="${ys(y)}" r="3.5" fill="var(--accent)"/>`);
 P.unk.forEach(([x,y])=>s+=`<rect x="${xs(x)-3.5}" y="${ys(y)-3.5}" width="7" height="7" fill="var(--ref)"/>`);
 if(!P.read.length)s+=`<text x="${(l+W)/2}" y="${H/2}" font-size="11" text-anchor="middle" fill="var(--muted)">Kurva terbentuk saat standar dibaca (Read Std)</text>`;
 g.innerHTML=s;}

/* --- mesin pemutaran --- */
const SIM={i:0,t:0,playing:false,speed:1,zoom:true,ff:false,fired:0,lang:'both',started:false};
const DIR={side:V(-1,.45,.35),console:V(.25,.75,1),comp:V(.2,1,.85),mono:V(.35,1,.6),lamp:V(.75,.85,.3),pd:V(-.7,.75,.7),all:V(.55,.6,.75)};
function cam(target,dist,dir){if(SIM.ff||!SIM.zoom)return;const p=typeof target==='string'?anchor(target):target.clone();flyTo(p,dist,DIR[dir]);}
function view(m){S.caseMode=m;S.explode=0;applyCase();syncBtns();}
function power(on){P.power=on;S.lampsOn=on;S.beams=on;syncBtns();if(!on){S.forceLamp=null;drawLCD(['','',' (daya mati)'],'#55605e');}}
function setWL(v){S.wl=v;$('wl').value=v;$('wlv').textContent=(+v).toFixed(1)+' nm';updateStat();}
function lcdSet(l){if(P.power)drawLCD(l);}
const CMP=V(-16,8.5,-1.8),LIDP=V(-16,20,0);
// urutan buka tutup → pasang kuvet → tutup tutup
function cuv(pairs,t0=0){return [[t0,()=>{view('cut');cam(LIDP,30,'comp');S.lidOpen=1}],[t0+.14,()=>pairs.forEach(([sl,k])=>setCuv(sl,k))],[t0+.26,()=>{S.lidOpen=0;log('Kuvet: '+pairs.map(([sl,k])=>`${sl} = ${solBy(k).n}`).join('; '))}]];}
function sweep(p0,p1){return p=>{if(p<p0)return;const u=Math.min(1,(p-p0)/(p1-p0));setWL(+(SCAN[0]+(SCAN[1]-SCAN[0])*u).toFixed(1));};}
const G={A:'A · Persiapan / Preparation',B:'B · Mencari λ maksimum / Finding λmax',Bp:'B · Print data / Printing',C:'B · Analisa sampel / Sample analysis',D:'C · Mematikan alat / Shutting down',E:'D · Mengeringkan kuvet / Drying cuvettes'};
const SSTEPS=[
 {g:'A',id:'Hubungkan kabel terminal utama ke sumber listrik.',en:'Connect the main power cable to the mains supply.',dur:4,
  acts:[[0,()=>{view('solid');cam(V(-27.5,9,-9),40,'side')}],[.5,()=>log('Kabel terminal utama terhubung.')]]},
 {g:'A',id:'Tekan tombol power pada alat UPS.',en:'Press the power button on the UPS.',dur:3.5,note:'UPS tidak dimodelkan. / The UPS is not modelled.',
  acts:[[0,()=>cam(V(0,8,0),85,'all')],[.5,()=>log('UPS menyala.')]]},
 {g:'A',id:'Keluarkan silika gel dari compartment.',en:'Remove the silica gel from the sample compartment.',dur:5,
  acts:[[0,()=>{view('cut');cam(LIDP,30,'comp')}],[.12,()=>S.lidOpen=1],[.45,()=>{silica.visible=false;log('Silika gel dikeluarkan.')}],[.75,()=>S.lidOpen=0]]},
 {g:'A',id:'Buka monitor alat UV-1700, lalu nyalakan alat dengan menekan tombol power di bagian kiri alat.',en:'Raise the UV-1700 display, then switch the instrument on with the power switch on its left side.',dur:5,
  acts:[[0,()=>{view('solid');cam('con',32,'console')}],[.35,()=>cam('psw',22,'side')],[.65,()=>{power(true);S.beams=false;syncBtns();drawLCD(['Initialize','','']);log('Alat dinyalakan.')}]]},
 {g:'A',id:'Tunggu sampai alat UV-1700 selesai melakukan pengecekan otomatis.',en:'Wait until the UV-1700 completes its automatic self-check.',dur:11,
  note:'Butir pengecekan menurut Service Manual §6.1/§9.4; urutan tampilan pada layar dapat berbeda. / Check items follow Service Manual §6.1/§9.4; on-screen order may differ.',
  acts:[[0,()=>{view('cut');cam(V(10,8.5,-12),40,'mono')}],
   [.06,()=>{lcdSet(['Initialize','ROM Check','OK']);log('ROM Check … OK')}],[.12,()=>{lcdSet(['Initialize','RAM Check','OK']);log('RAM Check … OK')}],
   [.2,()=>{cam('f',18,'mono');lcdSet(['Initialize','Filter Origin','']);log('Filter Origin … OK')}],
   [.34,()=>{cam('m1',20,'lamp');S.beams=true;syncBtns();setWL(600);lcdSet(['Initialize','Light Source Org.','']);log('Light Source Org. … OK')}],
   [.48,()=>{setWL(600);lcdSet(['Initialize','W Lamp Energy','']);log('W Lamp Energy … OK')}],
   [.6,()=>{setWL(250);lcdSet(['Initialize','D2 Lamp Energy','']);log('D2 Lamp Energy … OK')}],
   [.74,()=>{cam('g',22,'mono');S.forceLamp='d2';setWL(656.1);lcdSet(['Initialize','λ Check 656.1nm','']);log('λ Check (garis D₂ 656,1 nm) … OK')}],
   [.97,()=>{S.forceLamp=null;setWL(550);lcdSet(['Inisialisasi','selesai',''])}]]},
 {g:'A',id:'Setelah pengecekan selesai, tunggu selama 15 menit (warming up); setelah itu alat dapat digunakan.',en:'When the check is complete, allow a 15-minute warm-up; the instrument is then ready for use.',dur:7,
  note:'Waktu dipercepat. / Time is accelerated.',
  acts:[[0,()=>{view('cut');cam(V(6,8,-6),70,'all')}],[1,()=>log('Pemanasan 15 menit selesai.')]],
  tick:p=>lcdSet(['Warming up',`${Math.ceil(15*(1-p))} menit`,'550.0 nm'])},
 {g:'A',id:'Tekan F4 (PC Control) pada alat UV-1700.',en:'Press F4 (PC Control) on the UV-1700.',dur:4,
  acts:[[0,()=>{view('solid');cam('con',26,'console')}],[.5,()=>{lcdSet(['PC Control','','']);log('Mode PC Control aktif.')}]]},
 {g:'A',id:'Tekan tombol power pada CPU, monitor komputer, dan printer. Klik 2× UVProbe → Connect.',en:'Switch on the PC, monitor and printer. Double-click UVProbe → Connect.',dur:4.5,note:'Komputer tidak dimodelkan. / The PC is not modelled.',
  acts:[[0,()=>cam(V(0,8,0),85,'all')],[.6,()=>log('UVProbe: Connect.')]]},
 {g:'B',id:'Klik Window → Spectrum. Klik Edit → Method: isi wavelength range (start: λ tertinggi; end: λ terendah), pilih scan speed, klik Instrument Parameters, pilih measuring mode Absorbance, klik OK.',
  en:'Click Window → Spectrum. Click Edit → Method: enter the wavelength range (start: highest λ; end: lowest λ), choose the scan speed, click Instrument Parameters, select measuring mode Absorbance, click OK.',dur:6,
  acts:[[0,()=>cam(V(0,8,0),80,'all')],[.6,()=>log(`Method Spectrum: ${SCAN[0]}–${SCAN[1]} nm, mode Absorbance.`)]]},
 {g:'B',id:'Masukkan blanko pada kedua compartment → klik Baseline → OK.',en:'Place the blank in both cell positions → click Baseline → OK.',dur:10,
  acts:[...cuv([['S','blk'],['R','blk']]),[.32,()=>cam('g',22,'mono')],[.66,()=>cam(CMP,26,'comp')],[1,()=>{setWL(SCAN[0]);log('Baseline terkoreksi pada rentang scan.')}]],tick:sweep(.32,.96)},
 {g:'B',id:'Ganti blanko pada sample compartment dengan standar konsentrasi tertinggi → klik Start → OK → tunggu sampai selesai.',en:'Replace the blank in the sample position with the highest-concentration standard → click Start → OK → wait until the scan finishes.',dur:11,
  acts:[...cuv([['S','std5']]),[.3,()=>cam('g',20,'mono')],[.55,()=>cam('f',18,'mono')],[.72,()=>cam('cellS',20,'comp')],[.86,()=>cam('pdS',18,'pd')],[1,()=>log('Scan spektrum standar tertinggi selesai.')]],tick:sweep(.3,.96)},
 {g:'B',id:'Klik Operation → Peak Pick untuk melihat λ maksimum.',en:'Click Operation → Peak Pick to read λmax.',dur:5,
  note:'λ maks tidak ada dalam data; nilai bawaan 540 nm adalah literatur kompleks Cr(VI)–DPC (EPA 7196A). Ganti di kolom λ maks. / λmax is not in the data; the default 540 nm is the literature value for the Cr(VI)–DPC complex. Change it in the λmax field.',
  acts:[[0,()=>cam('g',22,'mono')],[.4,()=>{setWL(WL);lcdSet(['PC Control',`${WL.toFixed(1)} nm`,'']);log(`Peak Pick: λ maks = ${WL} nm (isian operator).`)}]]},
 {g:'B',id:'Klik File → Save As → beri nama spektrum pada kolom file dengan format Tanggal_Kelas_Kelompok → Save.',en:'Click File → Save As → name the spectrum file Date_Class_Group → Save.',dur:4,
  acts:[[0,()=>cam(V(0,8,0),80,'all')],[.5,()=>log('Spektrum disimpan.')]]},
 {g:'Bp',id:'Klik 1× pada tabel peak pick → klik kanan → Print.',en:'Click once on the peak-pick table → right-click → Print.',dur:3.5,acts:[[.5,()=>log('Tabel peak pick dicetak.')]]},
 {g:'C',id:'Klik Window → Photometric. Klik Edit → Method. Pada kolom wavelength ketik panjang gelombang yang diinginkan → Add → Next.',en:'Click Window → Photometric. Click Edit → Method. Type the required wavelength in the wavelength field → Add → Next.',dur:6,
  acts:[[0,()=>{view('cut');cam('g',20,'mono')}],[.35,()=>{setWL(WL);log(`Method Photometric: λ ${WL} nm.`)}]]},
 {g:'C',id:'Pada kolom type pilih Multi Point; pada kolom unit ketik satuan; pada kolom WL1 pilih panjang gelombang; klik Zero Interception → Next 3×.',en:'Under Type choose Multi Point; type the unit; under WL1 choose the wavelength; click Zero Interception → Next ×3.',dur:5,
  acts:[[0,()=>cam(V(0,8,0),80,'all')],[.5,()=>log('Type Multi Point, unit mg/L, WL1 dipilih.')]]},
 {g:'C',id:'Pada kolom file name klik titik tiga (…) → pilih folder → beri nama Tanggal_Kelas_Kelompok → Open → Finish.',en:'In File name click the ellipsis (…) → choose a folder → name it Date_Class_Group → Open → Finish.',dur:4,acts:[[.5,()=>log('Berkas photometric dibuat.')]]},
 {g:'C',id:'Pada standard table, ketik nama standar pada kolom Sample ID, lalu ketik konsentrasi standar pada kolom Conc.',en:'In the standard table, type each standard name under Sample ID and its concentration under Conc.',dur:5,
  acts:[[.5,()=>log(`Standard table: ${STD.length} standar diisi.`)]]},
 {g:'C',id:'Ganti sample compartment dengan blanko background → klik Autozero.',en:'Place the background blank in the sample position → click Autozero.',dur:7,
  acts:[...cuv([['S','blk'],['R','blk']]),[.35,()=>cam(CMP,24,'comp')],[.65,()=>cam('pdS',18,'pd')],[.8,()=>{lcdSet(['PC Control',`${WL.toFixed(1)} nm`,'0.000 Abs']);log('Autozero: 0,000 Abs.')}]]},
 ...STD.map(([c,a],i)=>({g:'C',id:`Ganti sample compartment dengan standar ${c} mg/L → klik Read Std.`,en:`Place the ${c} mg/L standard in the sample position → click Read Std.`,dur:6,
  acts:[...cuv([['S','std'+i]]),[.34,()=>cam('cellS',20,'comp')],[.62,()=>cam('pdS',17,'pd')],[.78,()=>{P.read.push([c,a]);lcdSet(['PC Control',`${WL.toFixed(1)} nm`,`${a.toFixed(3)} Abs`]);log(`Read Std ${c} mg/L: A = ${a.toFixed(3)}`);drawCal();
    if(i===STD.length-1){const R=regress();log(`Kurva: A = ${R.b.toFixed(4)}C ${R.a<0?'−':'+'} ${Math.abs(R.a).toFixed(4)}; r = ${R.r.toFixed(4)}`);}}],[.9,()=>cam(CMP,34,'comp')]]})),
 {g:'C',id:'Klik Sample Table → beri nama sampel pada kolom Sample ID.',en:'Click Sample Table → enter the sample names under Sample ID.',dur:4,acts:[[.5,()=>log('Sample table diisi.')]]},
 ...['bs','u1','u2','u3'].map(k=>({g:'C',id:`Ganti sample compartment dengan ${solBy(k).n.toLowerCase()} → klik Read Unk.`,en:`Place ${({bs:'the sample blank',u1:'sample 1',u2:'sample 2',u3:'sample 3 + spike'})[k]} in the sample position → click Read Unk.`,dur:6,
  acts:[...cuv([['S',k]]),[.34,()=>cam('cellS',20,'comp')],[.62,()=>cam('pdS',17,'pd')],[.78,()=>{const s=solBy(k),R=regress();const c=(s.a-R.a)/R.b;P.unk.push([Math.max(0,c),Math.max(0,s.a)]);
   lcdSet(['PC Control',`${WL.toFixed(1)} nm`,`${s.a.toFixed(3)} Abs`]);log(`Read Unk ${s.n}: A = ${s.a.toFixed(3)} → C = ${c.toFixed(4)} mg/L`);drawCal();}],[.9,()=>cam(CMP,34,'comp')]]})),
 {g:'C',id:'Setelah selesai semua pengukuran klik File → Save.',en:'After all measurements, click File → Save.',dur:3.5,acts:[[0,()=>cam(V(0,8,0),80,'all')],[.5,()=>log('Data photometric disimpan.')]]},
 {g:'C',id:'Print kurva kalibrasi: klik kanan pada standard table → Print. Print data sampel: klik kanan pada sample table → Print.',en:'Print the calibration curve: right-click the standard table → Print. Print sample data: right-click the sample table → Print.',dur:4.5,acts:[[.5,()=>log('Kurva kalibrasi dan data sampel dicetak.')]]},
 {g:'D',id:'Catat semua kejadian pada Log Book Alat dan tanda tangani dengan nama jelas.',en:'Record all events in the instrument log book and sign with your full name.',dur:4,acts:[[.5,()=>log('Log book diisi.')]]},
 {g:'D',id:'Klik Disconnect. Klik File → Exit.',en:'Click Disconnect. Click File → Exit.',dur:3.5,acts:[[.5,()=>log('UVProbe terputus dan ditutup.')]]},
 {g:'D',id:'Klik menu Windows → Shutdown Computer. Tekan tombol power monitor komputer dan printer.',en:'Open the Windows menu → Shut down the computer. Switch off the monitor and printer.',dur:4,acts:[[.5,()=>log('Komputer, monitor, dan printer dimatikan.')]]},
 {g:'D',id:'Tekan MODE pada alat UV-1700.',en:'Press MODE on the UV-1700.',dur:3.5,acts:[[0,()=>{view('solid');cam('con',26,'console')}],[.5,()=>{lcdSet(['Mode','','']);log('MODE ditekan.')}]]},
 {g:'D',id:'Tekan tombol power di bagian kiri untuk mematikan alat, tutup monitor alat, dan masukkan kembali silika gel ke dalam compartment.',en:'Switch the instrument off with the power switch on the left, fold down its display, and return the silica gel to the compartment.',dur:7,
  note:'Pengeluaran kuvet tersirat dalam langkah ini. / Removing the cuvettes is implied by this step.',
  acts:[[0,()=>cam('psw',22,'side')],[.25,()=>{power(false);log('Alat dimatikan.')}],[.4,()=>{view('cut');cam(LIDP,30,'comp');S.lidOpen=1}],[.55,()=>{setCuv('S',null);setCuv('R',null);silica.visible=true;log('Kuvet dikeluarkan; silika gel dimasukkan.')}],[.8,()=>S.lidOpen=0]]},
 {g:'D',id:'Tekan tombol OFF pada UPS. Cabut kabel dari kabel perpanjangan, kecuali kabel UPS.',en:'Press OFF on the UPS. Unplug the cables from the extension strip, except the UPS cable.',dur:4,acts:[[0,()=>{view('solid');cam(V(0,8,0),85,'all')}],[.5,()=>log('UPS dimatikan; kabel dicabut kecuali kabel UPS.')]]},
 {g:'D',id:'Rapikan area kerja di sekitar instrumen dan komputer dari sampel dan standar; rapikan kursi ke tempat aman.',en:'Clear samples and standards from around the instrument and computer; return the chairs to a safe place.',dur:4,acts:[[.5,()=>log('Area kerja dirapikan.')]]},
 {g:'E',id:'Setelah praktik selesai dan alat mati, keringkan kuvet dengan membilas etanol kemudian aseton. Balik kuvet (posisi aman) dan tunggu hingga kering sempurna.',en:'After the session, with the instrument off, dry the cuvettes by rinsing with ethanol and then acetone. Stand them inverted in a safe place until completely dry.',dur:5,
  note:'Khusus kuvet kuarsa dan kaca. / Quartz and glass cuvettes only.',acts:[[.5,()=>log('Kuvet dibilas etanol → aseton, dikeringkan terbalik.')]]}];

function resetSim(){P.S=P.R=null;P.read=[];P.unk=[];setCuv('S',null);setCuv('R',null);silica.visible=true;$('log').textContent='';drawCal();
 S.lidOpen=0;S.forceLamp=null;S.step=false;power(false);setWL(550);}
function fireUntil(s,p){while(SIM.fired<s.acts.length&&s.acts[SIM.fired][0]<=p){s.acts[SIM.fired][1]();SIM.fired++;}}
function startStep(i){SIM.i=Math.max(0,Math.min(SSTEPS.length-1,i));SIM.t=0;SIM.fired=0;fireUntil(SSTEPS[SIM.i],0);renderSim();}
function jump(i){i=Math.max(0,Math.min(SSTEPS.length-1,i));resetSim();SIM.ff=true;
 for(let k=0;k<i;k++){const s=SSTEPS[k];SIM.fired=0;fireUntil(s,1);s.tick&&s.tick(1);}
 SIM.ff=false;startStep(i);}
function simTick(dt){if(!SIM.playing)return;const s=SSTEPS[SIM.i];SIM.t+=dt*SIM.speed;const p=Math.min(1,SIM.t/s.dur);
 fireUntil(s,p);s.tick&&s.tick(p);$('simbar').style.width=(p*100).toFixed(1)+'%';
 if(p>=1){if(SIM.i<SSTEPS.length-1)startStep(SIM.i+1);else{SIM.playing=false;renderSim();}}}
const MENU=/(Window|Spectrum|Edit|Method|Instrument Parameters|Absorbance|Baseline|Start|OK|Operation|Peak Pick|File|Save As|Save|Print|Photometric|Add|Next|Multi Point|WL1|Zero Interception|Open|Finish|Autozero|Read Std|Read Unk|Sample Table|Sample ID|Conc|Connect|Disconnect|Exit|Shutdown Computer|Shut down|F4|MODE|UVProbe)(?![a-z])/g;
const chip=t=>t.replace(MENU,'<span class="menu">$1</span>');
function txt(s){const id=`<span lang="id">${chip(s.id)}</span>`,en=`<span lang="en">${chip(s.en)}</span>`;
 return SIM.lang==='id'?id:SIM.lang==='en'?en:`${id}<br><i style="opacity:.8">${en}</i>`;}
function renderSim(){const s=SSTEPS[SIM.i];
 $('simplay').textContent=SIM.playing?'⏸ Jeda':(SIM.t>0||SIM.i>0)?'▶ Lanjut':'▶ Putar';
 $('simplay').setAttribute('aria-pressed',SIM.playing);
 $('simcur').innerHTML=`<div class="eyebrow">${SIM.lang==='en'?'Step':'Langkah'} ${SIM.i+1} ${SIM.lang==='en'?'of':'dari'} ${SSTEPS.length} · ${G[s.g]}</div><p style="margin:4px 0">${txt(s)}</p>${s.note?`<div class="note" style="font-size:12.5px">${s.note}</div>`:''}`;
 const sc=$('simcap');sc.style.display='block';$('hint').hidden=true;
 sc.innerHTML=`<b>${SIM.i+1}/${SSTEPS.length}</b> · ${txt(s)}`;
 document.querySelectorAll('#steplist li').forEach((li,k)=>{li.classList.toggle('done',k<SIM.i);li.classList.toggle('cur',k===SIM.i);if(k===SIM.i&&SIM.playing)li.scrollIntoView({block:'nearest'});});
 $('simbar').style.width='0%';}
$('steplist').innerHTML=SSTEPS.map((s,i)=>`<li data-i="${i}" style="cursor:pointer"><span class="s">${String(i+1).padStart(2,'0')}</span><span><span class="eyebrow">${G[s.g].split(' / ')[0]}</span> ${s.id}<br><i style="color:var(--muted)">${s.en}</i></span></li>`).join('');
document.querySelectorAll('#steplist li').forEach(li=>li.onclick=()=>jump(+li.dataset.i));
$('simplay').onclick=()=>{if(!SIM.started){SIM.started=true;jump(0);}else if(SIM.i===SSTEPS.length-1&&SIM.t>=SSTEPS[SIM.i].dur){jump(0);}SIM.playing=!SIM.playing;renderSim();};
$('simprev').onclick=()=>{SIM.started=true;jump(SIM.i-1);};
$('simnext').onclick=()=>{SIM.started=true;jump(SIM.i+1);};
$('simrst').onclick=()=>{SIM.started=true;jump(0);};
document.querySelectorAll('[data-speed]').forEach(b=>b.onclick=()=>{SIM.speed=+b.dataset.speed;document.querySelectorAll('[data-speed]').forEach(x=>x.setAttribute('aria-pressed',x===b));});
$('simlang').onchange=e=>{SIM.lang=e.target.value;renderSim();};
$('simzoom').onchange=e=>{SIM.zoom=e.target.checked;};
$('simwl').onchange=e=>{const v=+e.target.value;if(v>=190&&v<=1100)WL=v;};
drawCal();setCuv('S',null);setCuv('R',null);
// putar otomatis saat tab praktikum pertama kali dibuka
document.querySelector('[data-t="tPrak"]').addEventListener('click',()=>{if(!SIM.started){SIM.started=true;jump(0);SIM.playing=true;renderSim();}});

/* ---------- data tab ---------- */
(function(){const x=STD.map(s=>s[0]),y=STD.map(s=>s[1]);const n=6,mx=x.reduce((a,b)=>a+b)/n,my=y.reduce((a,b)=>a+b)/n;let sxx=0,sxy=0,syy=0;x.forEach((v,i)=>{sxx+=(v-mx)**2;sxy+=(v-mx)*(y[i]-my);syy+=(y[i]-my)**2});
 const b=sxy/sxx,a=my-b*mx,r=sxy/Math.sqrt(sxx*syy);const res=y.map((v,i)=>v-(a+b*x[i]));const syx=Math.sqrt(res.reduce((s,v)=>s+v*v,0)/(n-2));
 const kad=A=>(A-a)/b*100*0.25/0.0004;const k1=kad(.191),k2=kad(.184),k3=kad(.193);
 const orig={k1:15397.46,k2:14822.19,k3:15495.02,c2:2485.71};const C1=(orig.k1+orig.k2)/2,rpd=Math.abs(orig.k1-orig.k2)/C1*100,rec=(orig.k3-C1)/orig.c2*100;
 const C1d=(k1+k2)/2,rpdd=Math.abs(k1-k2)/C1d*100,c2d=1000*0.001/0.0004,recd=(k3-C1d)/c2d*100;
 const sd=Math.abs(orig.k1-orig.k2)/Math.SQRT2,uRec=Math.sqrt(sd*sd+sd*sd/2)/orig.c2*100;
 const f=(v,d=2)=>v.toLocaleString('id-ID',{minimumFractionDigits:d,maximumFractionDigits:d});
 $('dataBox').innerHTML=`
 <h3>Kurva kalibrasi (6 titik, 0–0,5 mg/L)</h3>
 <dl class="kv"><dt>Slope b</dt><dd>${b.toFixed(4)} L·mg⁻¹</dd><dt>Intersep a</dt><dd>${a.toFixed(4)}</dd><dt>r / R²</dt><dd>${r.toFixed(4)} / ${(r*r).toFixed(4)}</dd><dt>s<sub>y/x</sub></dt><dd>${syx.toFixed(4)} Abs</dd><dt>LOD / LOQ (3,3s/b; 10s/b)</dt><dd>${(3.3*syx/b).toFixed(3)} / ${(10*syx/b).toFixed(3)} mg/L</dd></dl>
 <p>r = 0,9998 memenuhi ambang linearitas umum (r ≥ 0,995), tetapi r bukan uji linearitas. Residual berpola (−, +, −, +, +, −) dengan simpangan terbesar pada titik 0,5 mg/L (−0,0046 Abs), indikasi kelengkungan negatif ringan di ujung atas. Absorbansi sampel (0,184–0,193) berada di tengah rentang kalibrasi, sehingga dampaknya kecil.</p>
 <h3 style="margin-top:10px">Presisi duplo</h3>
 <div class="eq">%RPD = |${f(orig.k1)} − ${f(orig.k2)}| / ${f(C1)} × 100 = ${rpd.toFixed(2)}%</div>
 <p><span class="pill v">≤ 20%</span>Memenuhi batas RPD duplikat Cr(VI) 20% (pedoman validasi data EPA untuk hasil ≥ 5×QL). Batas internal laboratorium <b>perlu dikonfirmasi</b>.</p>
 <h3 style="margin-top:10px">Akurasi (spike)</h3>
 <div class="eq">%Recovery = (${f(orig.k3)} − ${f(C1)}) / ${f(orig.c2)} × 100 = ${rec.toFixed(2)}%</div>
 <p><span class="pill x">di luar 75–125%</span>Aritmetika lembar Anda benar dan kesimpulan TMS konsisten dengan batas matrix spike EPA (75–125%) maupun AOAC (±92–105% pada tingkat ~1%). Acuan yang dipakai laboratorium <b>perlu dikonfirmasi</b>.</p>
 <div class="note"><b>Catatan telaah.</b><ol>
 <li>Kadar asli (~15 110 mg/kg) 6,1× lebih tinggi dari jumlah spike (2 486 mg/kg). Pedoman EPA tidak mengevaluasi recovery bila kadar sampel ≥ 4× spike, karena ketidakpastian selisih C3 − C1 mendominasi. Dari presisi duplo, u(%Rec) ≈ ${uRec.toFixed(0)} poin persen (k = 1); recovery 15,5% tetap ±4u di bawah 100%, sehingga indikasi kehilangan Cr(VI) nyata, tetapi rancangan spike perlu diperbaiki.</li>
 <li>Kenaikan absorbansi yang diharapkan dari spike 1 mg Cr dalam 250 mL dan FP 100 adalah 0,04 mg/L × ${b.toFixed(4)} = 0,031 Abs; yang teramati hanya ±0,006 Abs. Penurunan ini konsisten dengan reduksi Cr(VI) → Cr(III) oleh matriks (Fe²⁺, sulfida, bahan organik) atau kondisi preparasi asam; EPA 3060A merekomendasikan digesti alkali dan karakterisasi sifat reduktif matriks. Ini hipotesis yang perlu diuji, misalnya dengan spike pasca-preparasi.</li>
 <li>Bobot sampel ditampilkan 0,0004 kg, padahal nilai kadar dan C2 (2 485,71 = 1/0,0004023) menunjukkan bobot aktual memiliki desimal tersembunyi. Dari nilai tampilan diperoleh kadar ${f(k1)}, ${f(k2)}, ${f(k3)} mg/kg, RPD ${rpdd.toFixed(2)}% dan recovery ${recd.toFixed(1)}%; kesimpulan tidak berubah, tetapi berkas Excel asli diperlukan agar perhitungan dapat ditelusuri penuh.</li>
 <li>IK memuat opsi <span class="menu">Zero Interception</span> pada metode Photometric, sedangkan lembar Anda memakai intersep 0,0024. Pastikan konsentrasi yang dilaporkan dihitung dari persamaan yang sama.</li>
 <li>Kolom "Volume (L) = 0,05" tidak dipakai dalam rumus kadar dan perannya terhadap FP 100 perlu dijelaskan dalam metodologi.</li></ol></div>`;})();

/* ---------- sources ---------- */
$('srcs').innerHTML=[
 ['Shimadzu Corporation. <i>UV-Vis Spectrophotometer UV-1700 series Service Manual</i>, P/N 206-98366A. §2.2–2.4, §5, §6, §8.3, §9.1–9.4 (Fig. 9-1, 9-2), §10.3–10.12.','https://www.manualslib.com/manual/1225472/Shimadzu-Uv-1700-Series.html?page=91'],
 ['Shimadzu Corporation. <i>Instruction Manual (Operation) UV-1700</i>, P/N 206-94785 (arsip; tidak dibaca langsung, dicantumkan sebagai acuan resmi pengoperasian).','https://archive.org/details/uoftmanuals005'],
 ['Spesifikasi UV-1700 (lebar pita 1 nm, 190–1100 nm, pergantian lampu 295–364 nm, cahaya liar 0,04%). speciation.net Instrument Database.','https://speciation.net/Appl/Techniques/technique.html?id=1297'],
 ['Politeknik AKA Bogor. <i>Instruksi Kerja Penggunaan Alat Spektrofotometer UV-Vis Shimadzu UV 1700 PharmaSpec</i>, IK/LAB.INST/6, berlaku 31 Agustus 2026, Revisi –.',null],
 ['Foto alat dan komputer UV-1700 di Laboratorium Instrumentasi (29 September 2026): label model CAT. NO. 206-55401-34, sakelar daya dan kipas sisi kiri, port I/O-2 dan Centronics, kompartemen sampel.',null],
 ['US EPA. <i>Method 3060A: Alkaline Digestion for Hexavalent Chromium</i>, SW-846, Rev. 1, 1996.','https://19january2021snapshot.epa.gov/sites/static/files/2015-12/documents/3060a.pdf'],
 ['Connecticut DEEP. <i>QA/QC Requirements: Hexavalent Chromium by SW-846 Method 7196</i> (kompleks DPC, 540 nm).','https://portal.ct.gov/-/media/deep/site_clean_up/guidance/rcp/rcpmethod7196hexchromepdf.pdf'],
 ['Ohio EPA. <i>Checklist 9: Hexavalent Chromium Data Validation</i> (RPD ≤ 20%; MS 75–125%).','https://dam.assets.ohio.gov/raw/upload/epa.ohio.gov/Portals/30/RCRA/docs/Checklist%209%20-%20Hexavalent%20Chromium%20Data%20Validation.docx'],
 ['AAFCO. Draft method needs template, mengutip batas recovery AOAC Single-Laboratory Validation.','https://www.aafco.org/wp-content/uploads/2023/01/12_draft_method_needs_template.pdf'],
 ['Thompson, M., Ellison, S.L.R., Fajgelj, A., Willetts, P. (1999). Harmonised guidelines for the use of recovery information in analytical measurement. <i>Pure Appl. Chem.</i> 71, 337–348.','https://www.farm.ucl.ac.be/tpao/temp/Portail_stat/Cours_stat/DES_INDU/Qualite/Documents/recovery.pdf']
].map(([t,u])=>`<li>${t}${u?` <a href="${u}" target="_blank" rel="noopener">tautan</a>`:''}</li>`).join('');
})();
</script>
"""

PAGE_HEAD = """<!doctype html><html lang="id"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1,viewport-fit=cover">
<style>[hidden]{display:none!important}body{margin:0}img{max-width:100%}</style></head><body>"""


def html_with_data(inp: dict) -> str:
    """Sisipkan isian aktif ke simulasi 3D bila strukturnya cocok (6 standar, blanko, 2 duplo, 1 spike, mode absorbansi)."""
    h = HTML_3D
    std = [(c, a) for c, a in inp["std"] if _ok(c, a)]
    ok = (inp["mode"] == "abs" and len(std) == 6 and inp["blank"]["use"] and _ok(inp["blank"]["y"])
          and all(_ok(d.get("y")) for d in inp["dup"]) and _ok(inp["spk"].get("y")))
    if ok:
        arr = ",".join(f"[{c:g},{a:.4f}]" for c, a in std)
        h = re.sub(r"const STD=\[\[.*?\]\];", f"const STD=[{arr}];", h, count=1)
        rep = {"bs": ("Blanko sampel", inp["blank"]["y"]), "u1": ("Sampel 1", inp["dup"][0]["y"]),
               "u2": ("Sampel 2", inp["dup"][1]["y"]), "u3": ("Sampel 3 + spike", inp["spk"]["y"])}
        for k, (name, y) in rep.items():
            h = re.sub(r"\{k:'%s',n:'[^']*',a:-?[\d.]+" % k, "{k:'%s',n:'%s',a:%s" % (k, name, float(y)), h, count=1)
    # di aplikasi, perhitungan ada di menu Perhitungan; tab data contoh di halaman 3D disembunyikan
    h = h.replace('<button class="tab" role="tab" aria-selected="false" data-t="tData">Data Cr(VI)</button>', "")
    return PAGE_HEAD + h + "</body></html>"


# ---------------------------------------------------------------------------
# 5. ANTARMUKA STREAMLIT
# ---------------------------------------------------------------------------
MODE_OPT = ["Absorbansi (dihitung melalui kurva kalibrasi)", "Konsentrasi larutan ukur (sudah dari alat)"]
BASIS_OPT = ["Padat: kadar dalam mg/kg (memakai bobot dan volume labu)", "Cair: kadar dalam satuan konsentrasi (× FP)"]
NUM = dict(value=None, format="%.6f", placeholder="isi nilai")


def fmt(v, d=4):
    return "—" if not _ok(v) else f"{v:,.{d}f}"


def load_into_state(src: dict):
    ss = st.session_state
    for k in list(ss.keys()):
        if k.startswith(("in_", "std_", "d0_", "d1_", "s_", "e")) and k not in ("ed_crit",):
            del ss[k]
    ss["in_unit"] = src["unit"]
    ss["in_mode"] = MODE_OPT[0] if src["mode"] == "abs" else MODE_OPT[1]
    ss["in_basis"] = BASIS_OPT[0] if src["basis"] == "padat" else BASIS_OPT[1]
    ss["in_zero"] = src["zero"]
    ss["in_nstd"] = len(src["std"])
    for i, (c, a) in enumerate(src["std"]):
        ss[f"std_c{i}"], ss[f"std_a{i}"] = c, a
    ss["in_buse"], ss["in_by"], ss["in_bcorr"] = src["blank"]["use"], src["blank"]["y"], src["blank"]["corr"]
    for i, d in enumerate(src["dup"]):
        for f in ("kode", "m", "v", "fp", "y"):
            ss[f"d{i}_{f}"] = d.get(f)
    for f in ("kode", "m", "v", "fp", "y", "cstd", "vstd", "base"):
        ss[f"s_{f}"] = src["spk"].get(f)
    ss["in_nextra"] = len(src["extra"])


def init_state():
    ss = st.session_state
    if "in_unit" not in ss:
        load_into_state(EMPTY)
    ss.setdefault("crit", DEFAULT_CRIT.copy())


def sample_inputs(prefix: str, title: str, padat: bool, ylab: str, spike=False):
    st.markdown(f"**{title}**")
    cols = st.columns(5 if padat else 3)
    i = 0
    kode = cols[i].text_input("Kode", key=f"{prefix}_kode"); i += 1
    m = v = None
    if padat:
        m = cols[i].number_input("Bobot m (g)", key=f"{prefix}_m", min_value=0.0, **NUM); i += 1
        v = cols[i].number_input("V labu (mL)", key=f"{prefix}_v", min_value=0.0, **NUM); i += 1
    fp = cols[i].number_input("FP", key=f"{prefix}_fp", min_value=0.0, help="Faktor pengenceran = V akhir / V alikuot", **NUM); i += 1
    y = cols[i].number_input(ylab, key=f"{prefix}_y", **NUM)
    return {"kode": kode or title, "m": m, "v": v, "fp": fp, "y": y}


def gather() -> dict:
    """Bangun formulir isian manual dan kembalikan dict input."""
    ss = st.session_state
    a, b, c = st.columns([1, 2, 2])
    unit = a.text_input("Satuan konsentrasi", key="in_unit")
    mode = b.radio("Hasil ukur yang diisikan", MODE_OPT, key="in_mode")
    basis = c.radio("Jenis sampel", BASIS_OPT, key="in_basis")
    absmode, padat = mode == MODE_OPT[0], basis == BASIS_OPT[0]
    ylab = "Absorbansi" if absmode else f"C ukur ({unit})"
    std = []
    if absmode:
        st.markdown("#### 1. Deret standar")
        x1, x2 = st.columns([1, 3])
        n = x1.number_input("Jumlah standar", min_value=3, max_value=15, step=1, key="in_nstd")
        x2.checkbox("Kurva melalui titik nol (setara opsi *Zero Interception* di IK)", key="in_zero")
        grid = st.columns(3)
        for i in range(int(n)):
            g = grid[i % 3]
            cc = g.number_input(f"Standar {i + 1}: C ({unit})", key=f"std_c{i}", **NUM)
            aa = g.number_input(f"Standar {i + 1}: absorbansi", key=f"std_a{i}", **NUM)
            std.append((cc, aa))
    st.markdown("#### 2. Blanko sampel")
    b1, b2, b3 = st.columns(3)
    buse = b1.checkbox("Ada blanko sampel", key="in_buse")
    by = b2.number_input(f"Blanko: {ylab}", key="in_by", disabled=not buse, **NUM)
    bcorr = b3.checkbox("Kurangkan blanko dari sampel", key="in_bcorr", disabled=not buse)
    st.markdown("#### 3. Sampel duplo")
    dup = [sample_inputs("d0", "Duplo 1", padat, ylab), sample_inputs("d1", "Duplo 2", padat, ylab)]
    st.markdown("#### 4. Sampel + spike (penambahan standar)")
    spk = sample_inputs("s", "Sampel + spike", padat, ylab)
    q = st.columns(3)
    spk["cstd"] = q[0].number_input(f"C larutan standar spike ({unit})", key="s_cstd", min_value=0.0, **NUM)
    spk["vstd"] = q[1].number_input("V standar spike yang ditambahkan (mL)", key="s_vstd", min_value=0.0, **NUM)
    spk["base"] = q[2].number_input("Bobot sampel spike (g)" if padat else "V sampel/total setelah spike (mL)",
                                    key="s_base", min_value=0.0, **NUM)
    st.markdown("#### 5. Sampel lain (opsional)")
    ne = st.number_input("Jumlah sampel lain", min_value=0, max_value=20, step=1, key="in_nextra")
    extra = [sample_inputs(f"e{i}", f"Sampel lain {i + 1}", padat, ylab) for i in range(int(ne))]
    return {"unit": unit or "mg/L", "mode": "abs" if absmode else "conc", "basis": "padat" if padat else "cair",
            "zero": bool(ss.get("in_zero")), "std": std,
            "blank": {"use": buse, "y": by, "corr": bcorr and buse}, "dup": dup, "spk": spk, "extra": extra}


def page_3d():
    st.subheader("Model 3D, cara kerja, dan simulasi IK/LAB.INST/6")
    st.caption("Seret untuk memutar 360°, cubit/gulir untuk zoom, ketuk komponen untuk nama, fungsi, dan prinsip kerja. "
               "Simulasi praktikum memakai absorbansi dari isian menu Perhitungan bila tersedia 6 standar, blanko, "
               "2 duplo, dan 1 spike; bila tidak, simulasi memakai data contoh.")
    h = st.slider("Tinggi tampilan (px)", 900, 2200, 1500, 50)
    inp = st.session_state.get("last_inp", EXAMPLE)
    components.html(html_with_data(inp), height=h, scrolling=True)


def page_calc():
    ss = st.session_state
    st.subheader("Perhitungan (isian manual)")
    st.caption("Semua hasil dihitung ulang dari nilai mentah yang Anda isi. Hasil akhir pada Excel contoh tidak dipakai, "
               "sehingga kesalahan di lembar contoh tidak terbawa.")
    c1, c2, _ = st.columns([1, 1, 3])
    if c1.button("Isi contoh perhitungan"):
        load_into_state(EXAMPLE); st.rerun()
    if c2.button("Kosongkan semua"):
        load_into_state(EMPTY); st.rerun()
    inp = gather()
    ss["last_inp"] = inp
    st.divider()
    st.markdown("## Hasil")
    try:
        out = compute(inp)
    except ValueError as e:
        st.info(f"Kurva kalibrasi belum dapat dihitung: {e}")
        return
    U, K = inp["unit"], out["K"]
    c = out.get("calib")
    if c:
        st.markdown("### A. Kurva kalibrasi")
        m = st.columns(6)
        for col, (lab, v) in zip(m, [("Slope b", c.b), ("Intersep a", c.a), ("r", c.r), ("R²", c.r2),
                                     (f"LOD ({U})", c.lod), (f"LOQ ({U})", c.loq)]):
            col.metric(lab, fmt(v))
        st.latex(rf"A = {c.b:.4f}\,C {'+' if c.a >= 0 else '-'} {abs(c.a):.4f}\qquad r = {c.r:.4f}")
        res = out["res"]
        try:
            import altair as alt
            line = pd.DataFrame({"C": [0, c.xmax], "A": [c.a, c.a + c.b * c.xmax]})
            ch = (alt.Chart(res).mark_circle(size=70).encode(x=alt.X("C standar:Q", title=f"Konsentrasi ({U})"),
                                                              y=alt.Y("A terukur:Q", title="Absorbansi"))
                  + alt.Chart(line).mark_line().encode(x="C:Q", y="A:Q"))
            if not out["tab"].empty:
                pts = out["tab"].rename(columns={"C ukur": "C", "Hasil ukur": "A"})[["Kode", "C", "A"]]
                ch = ch + alt.Chart(pts).mark_square(size=60, color="#ea6a2e").encode(x="C:Q", y="A:Q", tooltip=["Kode", "A"])
            st.altair_chart(ch.properties(height=300))
        except ImportError:
            st.line_chart(res.set_index("C standar")[["A terukur", "A prediksi"]])
        st.dataframe(res.style.format("{:.4f}", na_rep="—"), hide_index=True)
        pat = "".join("+" if v > 0 else "−" for v in res["Residual"])
        st.caption(f"Pola tanda residual: {pat}. Pola sistematis menandakan penyimpangan linearitas yang tidak terlihat dari r.")
    tab = out["tab"]
    if tab.empty:
        st.info("Isi data sampel untuk menghitung kadar.")
        return
    st.markdown("### B. Kadar")
    num = [col for col in tab.columns if col not in ("Kode", "Jenis")]
    st.dataframe(tab.style.format({k: "{:,.4f}" for k in num}, na_rep="—"), hide_index=True)
    d0 = inp["dup"][0]
    if _ok(d0["y"]) and (c or inp["mode"] == "conc"):
        r0 = tab.iloc[0]
        with st.expander("Contoh substitusi (baris pertama)", expanded=True):
            if c:
                st.latex(rf"C=\frac{{A-a}}{{b}}=\frac{{{d0['y']}-({c.a:.4f})}}{{{c.b:.4f}}}={r0['C ukur']:.5f}\ \text{{{U}}}")
            if inp["blank"]["corr"]:
                st.latex(rf"C_{{terkoreksi}}=C-C_{{blanko}}={r0['C ukur']:.5f}-({out.get('c_blank', 0):.5f})={r0['C terkoreksi']:.5f}")
            if inp["basis"] == "padat":
                st.latex(rf"\text{{Kadar}}=\frac{{C\times FP\times V}}{{m}}=\frac{{{r0['C terkoreksi']:.5f}\times{r0['FP']:g}\times{r0['V (mL)']:g}}}{{{r0['m (g)']:g}}}={r0[f'Kadar ({K})']:.2f}\ \text{{{K}}}")
                st.caption("Satuan: (mg/L × mL) / g = mg/kg, karena 1 mL/g = 1 L/kg.")
            else:
                st.latex(rf"\text{{Kadar}}=C\times FP={r0['C terkoreksi']:.5f}\times{r0['FP']:g}={r0[f'Kadar ({K})']:.4f}\ \text{{{K}}}")
    st.markdown("### C. Presisi duplo")
    if "rpd" in out:
        st.latex(rf"\%RPD=\frac{{|x_1-x_2|}}{{(x_1+x_2)/2}}\times100=\frac{{|{out['x1']:.2f}-{out['x2']:.2f}|}}{{{out['c1']:.2f}}}\times100={out['rpd']:.2f}\%")
    else:
        st.info("Lengkapi kedua duplo.")
    st.markdown("### D. Akurasi (%recovery)")
    if "rec" in out:
        sp = inp["spk"]
        st.latex(rf"C_2=\frac{{C_{{std}}\,V_{{std}}}}{{{'m' if inp['basis'] == 'padat' else 'V'}}}=\frac{{{sp['cstd']:g}\times{sp['vstd']:g}}}{{{sp['base']:g}}}={out['c2']:.2f}\ \text{{{K}}}")
        st.latex(rf"\%Rec=\frac{{C_3-C_1}}{{C_2}}\times100=\frac{{{out['c3']:.2f}-{out['c1']:.2f}}}{{{out['c2']:.2f}}}\times100={out['rec']:.2f}\%")
        st.caption("C1 = rerata duplo, C3 = kadar sampel + spike; keduanya sudah dikoreksi pengenceran dan pada basis yang sama.")
        dg = out["diag"]
        st.dataframe(pd.DataFrame({"Diagnostik": list(dg), "Nilai": [round(float(v), 4) for v in dg.values()]}), hide_index=True)
    else:
        st.info("Lengkapi data sampel + spike dan larutan standar spike.")
    for w in out["warn"]:
        st.warning(w)
    st.markdown("### E. Evaluasi terhadap kriteria")
    ss.crit = st.data_editor(ss.crit, key="ed_crit", disabled=["Parameter", "Acuan"], hide_index=True)
    st.dataframe(evaluate(out, ss.crit), hide_index=True)
    st.caption("Status lab tetap 'Perlu dikonfirmasi' sampai batas internal laboratorium/IK diisi.")
    if "rec" in out:
        st.download_button("Unduh Excel berformula", data=build_workbook(inp), file_name="Perhitungan_UV1700.xlsx",
                           mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    else:
        st.caption("Tombol unduh Excel muncul setelah duplo dan spike lengkap.")


def page_src():
    st.subheader("Sumber dan status verifikasi")
    st.markdown("""
- Shimadzu Corporation, *UV-1700 series Service Manual* 206-98366A — §9.1 Fig. 9-1/9-2 (jalur optik, filter), §9.4 (CPU, preamp), §10 (daftar suku cadang), §2, §5, §6.
- Shimadzu Corporation, *Instruction Manual (Operation)* 206-94785 — acuan resmi pengoperasian (arsip).
- Politeknik AKA Bogor, *IK/LAB.INST/6* Instruksi Kerja Spektrofotometer UV-Vis Shimadzu UV 1700 PharmaSpec, berlaku 31 Agustus 2026.
- US EPA Method 3060A; CT DEEP RCP Method 7196; Ohio EPA Checklist 9 (Cr(VI) data validation).
- AOAC SLV (dikutip AAFCO); Thompson et al. (1999) *Pure Appl. Chem.* 71, 337–348.

**Terverifikasi**: komponen, urutan jalur optik, rentang filter, titik pergantian lampu, tipe detektor.
**Konseptual**: bentuk eksterior (disederhanakan, tanpa logo), jarak/sudut komponen, posisi PCB, besar rotasi M1 dan kisi,
penempatan S di depan/R di belakang (cocokkan dengan label alat).
""")


def main():
    st.set_page_config(page_title="UV-1700 PharmaSpec 3D", layout="wide")
    init_state()
    st.title("Spektrofotometer UV-Vis UV-1700 PharmaSpec")
    page = st.sidebar.radio("Menu", ["Model 3D & simulasi IK", "Perhitungan", "Sumber"])
    st.sidebar.caption("Data aktif dipakai bersama oleh simulasi 3D dan perhitungan.")
    {"Model 3D & simulasi IK": page_3d, "Perhitungan": page_calc, "Sumber": page_src}[page]()


if __name__ == "__main__":
    main()
